"""缺陷保存视图、批处理、导入导出、度量与可配置工作流。"""
import asyncio
import csv
import hashlib
import io
import json
import re
from copy import deepcopy
from datetime import datetime, timedelta, timezone
from statistics import mean, median
from uuid import UUID, uuid4, uuid5
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from loguru import logger
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, or_, select
from sqlalchemy.exc import IntegrityError

from app.config import settings
from app.models.database import (
    Defect, DefectEvent, DefectJob, DefectSeverity, DefectType, DefectView,
    DefectWorkflow, DefectWorkflowRevision, ProjectDefectSettings, ProjectModule,
    Project, ProjectIteration, ProjectMember, User,
)
from app.modules.auth.dependencies import get_current_user
from app.modules.defects.permissions import context_for
from app.modules.defects.query import enrich_list, filter_conditions
from app.modules.defects.service import (
    append_event, assign_defect, create_defect, delete_or_restore, definition_for,
    edit_defect, json_value, snapshot, transition_defect,
)
from app.modules.defects.workflow import DEFAULT_WORKFLOW, ROLE_KEYS, STATES, TRANSITION_FIELDS, fingerprint, project_state
from app.schemas.defect_management import (
    BatchOperation, DefectSettingsInput, ExportRequest, ImportCommitRequest,
    SavedViewInput, SavedViewUpdate, WorkflowCreate, WorkflowDraftUpdate,
    WorkflowMigrationRequest,
)
from app.schemas.defects import AssignmentRequest, DefectCreate, DefectUpdate, TransitionRequest
from app.utils.database import get_db_session
from app.utils.storage import get_minio_client

router = APIRouter()
VIEW_FILTERS = {"q", "severity", "defect_type", "priority", "workflow_states", "assignee_id", "reporter_id",
                "verifier_id", "module_id", "iteration_id", "environment_profile_id", "label", "source", "view",
                "created_from", "created_to", "updated_from", "updated_to", "overdue", "watching"}
VIEW_COLUMNS = {"defect_key", "title", "workflow_state", "severity", "priority", "defect_type", "module_name",
                "assignee_name", "verifier_name", "reporter_name", "source", "created_at", "updated_at", "due_at"}
SORT_COLUMNS = {"created_at", "updated_at", "due_at", "priority", "severity", "defect_key"}
IMPORT_HEADERS = ["标题", "描述", "严重程度", "优先级", "缺陷类型", "复现步骤", "预期结果", "实际结果", "标签", "处理人用户名", "验证人用户名"]
HEADER_MAP = {"标题": "title", "title": "title", "描述": "description", "description": "description",
              "严重程度": "severity", "severity": "severity", "优先级": "priority", "priority": "priority",
              "缺陷类型": "defect_type", "defect_type": "defect_type", "复现步骤": "reproduce_steps",
              "reproduce_steps": "reproduce_steps", "预期结果": "expected_result", "expected_result": "expected_result",
              "实际结果": "actual_result", "actual_result": "actual_result", "标签": "labels", "labels": "labels",
              "处理人用户名": "assignee_username", "assignee_username": "assignee_username",
              "验证人用户名": "verifier_username", "verifier_username": "verifier_username"}


def success(data, message="success"):
    return {"code": 0, "data": data, "message": message}


def _view_dict(item):
    return {"id": str(item.id), "project_id": str(item.project_id), "owner_id": str(item.owner_id), "name": item.name,
            "visibility": item.visibility, "query": item.query_json, "columns": item.columns_json,
            "sort": item.sort_json, "row_version": item.row_version, "created_at": json_value(item.created_at),
            "updated_at": json_value(item.updated_at)}


def _validate_view(data):
    unknown = set(data.query) - VIEW_FILTERS
    if unknown:
        raise HTTPException(422, "视图包含不支持的筛选字段：" + "、".join(sorted(unknown)))
    if any(key not in VIEW_COLUMNS for key in data.columns):
        raise HTTPException(422, "视图包含不支持的列表字段")
    for item in data.sort:
        if set(item) != {"field", "direction"} or item["field"] not in SORT_COLUMNS or item["direction"] not in {"asc", "desc"}:
            raise HTTPException(422, "视图排序规则无效")


@router.get("/views")
async def list_views(project_id: UUID, db=Depends(get_db_session), user=Depends(get_current_user)):
    await context_for(db, user, project_id)
    rows = (await db.execute(select(DefectView).where(DefectView.project_id == project_id,
            or_(DefectView.owner_id == user.id, DefectView.visibility == "project")).order_by(DefectView.name))).scalars().all()
    return success({"items": [_view_dict(item) for item in rows]})


@router.post("/views")
async def create_view(project_id: UUID, req: SavedViewInput, db=Depends(get_db_session), user=Depends(get_current_user)):
    context = await context_for(db, user, project_id, write=True)
    if req.visibility == "project" and not context.privileged:
        raise HTTPException(403, "项目共享视图仅限项目负责人创建")
    _validate_view(req)
    item = DefectView(id=uuid4(), project_id=project_id, owner_id=user.id, name=req.name,
                      visibility=req.visibility, query_json=req.query, columns_json=req.columns, sort_json=req.sort)
    db.add(item)
    try:
        await db.commit()
    except IntegrityError:
        await db.rollback()
        raise HTTPException(409, "同名视图已存在")
    return success(_view_dict(item), "视图已保存")


@router.put("/views/{view_id}")
async def update_view(view_id: UUID, req: SavedViewUpdate, db=Depends(get_db_session), user=Depends(get_current_user)):
    item = (await db.execute(select(DefectView).where(DefectView.id == view_id).with_for_update())).scalar_one_or_none()
    if not item:
        raise HTTPException(404, "视图不存在")
    context = await context_for(db, user, item.project_id, write=True)
    if item.owner_id != user.id and not context.privileged:
        raise HTTPException(403, "无权修改该视图")
    if item.row_version != req.expected_version:
        raise HTTPException(409, "视图已被修改，请刷新后重试")
    if req.visibility == "project" and not context.privileged:
        raise HTTPException(403, "项目共享视图仅限项目负责人维护")
    _validate_view(req)
    item.name, item.visibility, item.query_json, item.columns_json, item.sort_json = req.name, req.visibility, req.query, req.columns, req.sort
    item.row_version += 1
    item.updated_at = datetime.now(timezone.utc)
    await db.commit()
    return success(_view_dict(item))


@router.delete("/views/{view_id}")
async def delete_view(view_id: UUID, db=Depends(get_db_session), user=Depends(get_current_user)):
    item = await db.get(DefectView, view_id)
    if not item:
        raise HTTPException(404, "视图不存在")
    context = await context_for(db, user, item.project_id, write=True)
    if item.owner_id != user.id and not context.privileged:
        raise HTTPException(403, "无权删除该视图")
    await db.delete(item)
    await db.commit()
    return success({"id": str(view_id)})


async def _batch_one(db, user, req, defect):
    values = req.values
    versions = values.get("expected_versions") or {}
    expected = versions.get(str(defect.id))
    if req.preview:
        return {"id": str(defect.id), "defect_key": defect.defect_key, "row_version": defect.row_version,
                "can_apply": expected in (None, defect.row_version)}
    if expected is None:
        raise HTTPException(422, "执行批量操作必须携带预览返回的版本号")
    if req.operation == "assign":
        payload = AssignmentRequest(expected_version=expected, request_id=uuid5(req.project_id, f"batch:assign:{defect.id}:{expected}"),
                                    assignee_id=values.get("assignee_id"), verifier_id=values.get("verifier_id"),
                                    note=values.get("note", "批量指派"))
        updated, _ = await assign_defect(db, user, defect.id, payload)
    elif req.operation == "transition":
        payload = TransitionRequest(action_id=values.get("action_id", ""), expected_version=expected,
            request_id=uuid5(req.project_id, f"batch:transition:{defect.id}:{expected}:{values.get('action_id')}"),
            note=values.get("note", "批量流转"), fields=values.get("fields") or {})
        updated, _ = await transition_defect(db, user, defect.id, payload)
    elif req.operation == "delete":
        updated = await delete_or_restore(db, user, defect.id, expected)
    else:
        changes = {}
        if req.operation == "priority":
            changes["priority"] = values.get("priority")
        elif req.operation == "module":
            changes["module_id"] = values.get("module_id")
        elif req.operation == "iteration":
            changes["iteration_id"] = values.get("iteration_id")
        elif req.operation in {"add_labels", "remove_labels"}:
            labels = set(defect.labels or [])
            incoming = {str(value).strip() for value in values.get("labels", []) if str(value).strip()}
            labels = labels | incoming if req.operation == "add_labels" else labels - incoming
            changes["labels"] = sorted(labels)
        else:
            raise HTTPException(422, "未知批量操作")
        updated = await edit_defect(
            db, user, defect.id, DefectUpdate(expected_version=expected, note=values.get("note", "批量更新"), **changes))
    return {"id": str(updated.id), "defect_key": updated.defect_key, "row_version": updated.row_version}


@router.post("/batch")
async def batch_operation(req: BatchOperation, db=Depends(get_db_session), user=Depends(get_current_user)):
    await context_for(db, user, req.project_id, write=True)
    found = (await db.execute(select(Defect).where(Defect.id.in_(req.defect_ids), Defect.project_id == req.project_id))).scalars().all()
    by_id = {item.id: item for item in found}
    succeeded, failed = [], []
    for identifier in req.defect_ids:
        item = by_id.get(identifier)
        if not item or item.deleted_at:
            failed.append({"id": str(identifier), "message": "缺陷不存在或已删除"})
            continue
        defect_key = item.defect_key
        try:
            async with db.begin_nested():
                if req.preview:
                    values = {**req.values, "expected_versions": {str(item.id): item.row_version}}
                    simulated = req.model_copy(update={"preview": False, "values": values})
                    await _batch_one(db, user, simulated, item)
                    await db.flush()
                    succeeded.append({"id": str(identifier), "defect_key": defect_key,
                                      "row_version": values["expected_versions"][str(identifier)], "can_apply": True})
                else:
                    succeeded.append(await _batch_one(db, user, req, item))
        except HTTPException as exc:
            detail = exc.detail.get("message") if isinstance(exc.detail, dict) else str(exc.detail)
            failed.append({"id": str(identifier), "defect_key": defect_key, "message": detail})
        except Exception:  # noqa: BLE001
            logger.exception("批量缺陷操作失败: project={} defect={} operation={}", req.project_id, identifier, req.operation)
            failed.append({"id": str(identifier), "defect_key": defect_key, "message": "内部处理失败，请稍后重试"})
    if req.preview:
        await db.rollback()
    else:
        await db.commit()
    return success({"succeeded": succeeded, "failed": failed, "preview": req.preview,
                    "summary": {"success": len(succeeded), "failed": len(failed)}})


def _safe_cell(value):
    text = "" if value is None else str(value)
    return "'" + text if text.startswith(("=", "+", "-", "@", "\t", "\r")) else text


def _workbook_bytes(rows):
    book = Workbook()
    sheet = book.active
    sheet.title = "缺陷"
    sheet.append(IMPORT_HEADERS)
    for row in rows:
        values = [row.get(key) for key in ("title", "description", "severity", "priority", "defect_type",
            "reproduce_steps", "expected_result", "actual_result", "labels", "assignee_name", "verifier_name")]
        values[5] = "\n".join(values[5] or []) if isinstance(values[5], list) else values[5]
        values[8] = ",".join(values[8] or []) if isinstance(values[8], list) else values[8]
        sheet.append([_safe_cell(value) for value in values])
    for width, column in zip((36, 55, 12, 12, 14, 55, 45, 45, 24, 18, 18), sheet.columns):
        sheet.column_dimensions[column[0].column_letter].width = width
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


@router.get("/import-template")
async def import_template(user=Depends(get_current_user)):
    data = _workbook_bytes([])
    return StreamingResponse(io.BytesIO(data), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=defect-import-template.xlsx"})


def _read_import(filename, data):
    if filename.lower().endswith(".csv"):
        text = data.decode("utf-8-sig")
        return list(csv.DictReader(io.StringIO(text)))
    if filename.lower().endswith(".xlsx"):
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = book.active
        values = list(sheet.iter_rows(values_only=True))
        if not values:
            return []
        headers = [str(value or "").strip() for value in values[0]]
        return [{headers[index]: value for index, value in enumerate(row) if index < len(headers)} for row in values[1:]]
    raise HTTPException(415, "仅支持 CSV 或 XLSX")


async def _normalize_import(db, project_id, raw_rows):
    project = await db.get(Project, project_id)
    member_ids = set((await db.execute(select(ProjectMember.user_id).where(ProjectMember.project_id == project_id,
        ProjectMember.access == "write"))).scalars().all())
    if project:
        member_ids.add(project.owner_id)
    users = {item.username: item for item in (await db.execute(select(User).where(
        User.id.in_(member_ids), User.is_active.is_(True)))).scalars().all()}
    existing_titles = {value.lower() for value in (await db.execute(select(Defect.title).where(
        Defect.project_id == project_id, Defect.deleted_at.is_(None)))).scalars().all()}
    severity_map = {"致命": "P0", "严重": "P1", "一般": "P2", "轻微": "P3"}
    priority_map = {"紧急": "urgent", "高": "high", "普通": "normal", "低": "low"}
    type_map = {"业务缺陷": "business", "程序缺陷": "program", "性能缺陷": "performance", "集成缺陷": "integration", "安全缺陷": "security"}
    result = []
    for number, raw in enumerate(raw_rows, start=2):
        row = {HEADER_MAP.get(str(key).strip().lower(), HEADER_MAP.get(str(key).strip(), str(key))): value for key, value in raw.items()}
        normalized = {"title": str(row.get("title") or "").strip(), "description": str(row.get("description") or "").strip(),
                      "severity": severity_map.get(str(row.get("severity") or "").strip(), str(row.get("severity") or "P2").strip()),
                      "priority": priority_map.get(str(row.get("priority") or "").strip(), str(row.get("priority") or "normal").strip()),
                      "defect_type": type_map.get(str(row.get("defect_type") or "").strip(), str(row.get("defect_type") or "business").strip()),
                      "reproduce_steps": [item.strip() for item in re.split(r"[\r\n]+|\s*;\s*", str(row.get("reproduce_steps") or "")) if item.strip()],
                      "expected_result": str(row.get("expected_result") or "").strip() or None,
                      "actual_result": str(row.get("actual_result") or "").strip() or None,
                      "labels": [item.strip() for item in re.split(r"[,，;；]", str(row.get("labels") or "")) if item.strip()]}
        errors, warnings = [], []
        if len(normalized["title"]) < 2:
            errors.append("标题至少 2 个字符")
        if not normalized["description"]:
            errors.append("描述不能为空")
        if normalized["severity"] not in {"P0", "P1", "P2", "P3"}:
            errors.append("严重程度无效")
        if normalized["priority"] not in {"urgent", "high", "normal", "low"}:
            errors.append("优先级无效")
        if normalized["defect_type"] not in {"business", "program", "performance", "integration", "security"}:
            errors.append("缺陷类型无效")
        for source, target in (("assignee_username", "assignee_id"), ("verifier_username", "verifier_id")):
            name = str(row.get(source) or "").strip()
            if name and name not in users:
                errors.append(f"用户 {name} 不存在或已停用")
            elif name:
                normalized[target] = str(users[name].id)
        if normalized["title"].lower() in existing_titles:
            warnings.append("项目中已有同标题缺陷")
        result.append({"row_number": number, "data": normalized, "errors": errors, "warnings": warnings, "valid": not errors})
    return result


@router.post("/imports/preview")
async def preview_import(project_id: UUID = Form(...), request_key: UUID = Form(...), file: UploadFile = File(...),
                         db=Depends(get_db_session), user=Depends(get_current_user)):
    await context_for(db, user, project_id, write=True)
    data = await file.read(20 * 1024 * 1024 + 1)
    if len(data) > 20 * 1024 * 1024:
        raise HTTPException(413, "导入文件不能超过 20MB")
    digest = hashlib.sha256(data).hexdigest()
    existing = (await db.execute(select(DefectJob).where(DefectJob.project_id == project_id,
        DefectJob.creator_id == user.id, DefectJob.kind == "import_preview", DefectJob.request_key == request_key))).scalar_one_or_none()
    if existing:
        if existing.input_hash != digest:
            raise HTTPException(409, "同一请求编号不能用于不同文件")
        return success({"job_id": str(existing.id), **(existing.summary_json or {})})
    raw = _read_import(file.filename or "", data)
    if len(raw) > 5000:
        raise HTTPException(422, "单次最多导入 5000 条缺陷")
    rows = await _normalize_import(db, project_id, raw)
    summary = {"filename": file.filename, "total": len(rows), "valid": sum(item["valid"] for item in rows),
               "invalid": sum(not item["valid"] for item in rows), "rows": rows}
    job = DefectJob(id=uuid4(), project_id=project_id, creator_id=user.id, kind="import_preview", status="completed",
                    request_key=request_key, input_hash=digest, summary_json=summary,
                    expires_at=datetime.now(timezone.utc) + timedelta(days=7), completed_at=datetime.now(timezone.utc))
    db.add(job)
    await db.commit()
    return success({"job_id": str(job.id), **summary})


@router.post("/imports/{job_id}/commit")
async def commit_import(job_id: UUID, req: ImportCommitRequest, db=Depends(get_db_session), user=Depends(get_current_user)):
    job = (await db.execute(select(DefectJob).where(DefectJob.id == job_id).with_for_update())).scalar_one_or_none()
    if not job or job.kind != "import_preview" or job.creator_id != user.id:
        raise HTTPException(404, "导入预览不存在")
    await context_for(db, user, job.project_id, write=True)
    if job.expires_at and job.expires_at < datetime.now(timezone.utc):
        raise HTTPException(410, "导入预览已过期")
    input_hash = fingerprint(req.model_dump(mode="json", exclude={"request_key"}))
    prior = (await db.execute(select(DefectJob).where(DefectJob.project_id == job.project_id,
        DefectJob.creator_id == user.id, DefectJob.kind == "import_commit", DefectJob.request_key == req.request_key))).scalar_one_or_none()
    if prior:
        if prior.input_hash != input_hash:
            raise HTTPException(409, "同一请求编号不能用于不同导入范围")
        return success({"job_id": str(prior.id), **(prior.summary_json or {})})
    selected = set(req.row_numbers or [])
    rows = [item for item in (job.summary_json or {}).get("rows", []) if item.get("valid") and (not selected or item["row_number"] in selected)]
    succeeded, failed = [], []
    for item in rows:
        try:
            async with db.begin_nested():
                data = dict(item["data"])
                data["project_id"] = job.project_id
                data["creation_request_id"] = uuid5(job.id, f"row:{item['row_number']}")
                defect = await create_defect(db, user, DefectCreate(**data))
                defect.source = "import"
                succeeded.append({"row_number": item["row_number"], "id": str(defect.id), "defect_key": defect.defect_key})
        except HTTPException as exc:
            detail = exc.detail.get("message") if isinstance(exc.detail, dict) else str(exc.detail)
            failed.append({"row_number": item["row_number"], "message": detail[:300]})
        except Exception:  # noqa: BLE001
            logger.exception("缺陷导入失败: project={} row={}", job.project_id, item["row_number"])
            failed.append({"row_number": item["row_number"], "message": "内部处理失败，请核对数据后重试"})
    summary = {"succeeded": succeeded, "failed": failed, "success": len(succeeded), "failure": len(failed)}
    commit_job = DefectJob(id=uuid4(), project_id=job.project_id, creator_id=user.id, kind="import_commit", status="completed",
                           request_key=req.request_key, input_hash=input_hash, summary_json=summary,
                           expires_at=datetime.now(timezone.utc) + timedelta(days=30), completed_at=datetime.now(timezone.utc))
    db.add(commit_job)
    await db.commit()
    return success({"job_id": str(commit_job.id), **summary}, "导入完成")


@router.post("/exports")
async def create_export(req: ExportRequest, db=Depends(get_db_session), user=Depends(get_current_user)):
    await context_for(db, user, req.project_id)
    digest = fingerprint(req.model_dump(mode="json", exclude={"request_key"}))
    prior = (await db.execute(select(DefectJob).where(DefectJob.project_id == req.project_id,
        DefectJob.creator_id == user.id, DefectJob.kind == "export", DefectJob.request_key == req.request_key))).scalar_one_or_none()
    if prior:
        if prior.input_hash != digest:
            raise HTTPException(409, "同一请求编号不能用于不同导出条件")
        return success(_job_dict(prior))
    query = select(Defect).where(Defect.project_id == req.project_id, Defect.deleted_at.is_(None))
    if req.defect_ids:
        query = query.where(Defect.id.in_(req.defect_ids))
    else:
        allowed = {"severity", "defect_type", "priority", "workflow_state", "assignee_id", "verifier_id", "module_id",
                   "iteration_id", "environment_profile_id", "label", "source", "q", "view", "watching", "overdue",
                   "created_from", "created_to", "updated_from", "updated_to"}
        unknown = set(req.filters) - allowed
        if unknown:
            raise HTTPException(422, "导出条件包含不支持的字段：" + "、".join(sorted(unknown)))
        values = dict(req.filters)
        if values.get("workflow_state"):
            values["workflow_states"] = [values.pop("workflow_state")]
        for key in ("assignee_id", "verifier_id", "module_id", "iteration_id", "environment_profile_id"):
            if values.get(key):
                try: values[key] = UUID(str(values[key]))
                except ValueError: raise HTTPException(422, f"无效的导出条件 {key}")
        for key in ("created_from", "created_to", "updated_from", "updated_to"):
            if values.get(key):
                try: values[key] = datetime.fromisoformat(str(values[key]).replace("Z", "+00:00"))
                except ValueError: raise HTTPException(422, f"无效的导出时间条件 {key}")
        query = query.where(*filter_conditions(user_id=user.id, **values))
    rows = (await db.execute(query.order_by(Defect.created_at.desc()).limit(10000))).scalars().all()
    enriched = await enrich_list(db, rows)
    if req.format == "xlsx":
        data, content_type, suffix = _workbook_bytes(enriched), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"
    else:
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(IMPORT_HEADERS)
        for row in enriched:
            values = [row.get(key) for key in ("title", "description", "severity", "priority", "defect_type",
                "reproduce_steps", "expected_result", "actual_result", "labels", "assignee_name", "verifier_name")]
            values[5] = "\n".join(values[5] or []) if isinstance(values[5], list) else values[5]
            values[8] = ",".join(values[8] or []) if isinstance(values[8], list) else values[8]
            writer.writerow([_safe_cell(value) for value in values])
        data, content_type, suffix = output.getvalue().encode("utf-8-sig"), "text/csv", "csv"
    identifier = uuid4()
    object_key = f"defect-jobs/{req.project_id}/{identifier}/defects.{suffix}"
    await asyncio.to_thread(get_minio_client().put_object, settings.MINIO_BUCKET, object_key, io.BytesIO(data), len(data), content_type=content_type)
    job = DefectJob(id=identifier, project_id=req.project_id, creator_id=user.id, kind="export", status="completed",
                    request_key=req.request_key, input_hash=digest, result_key=object_key,
                    summary_json={"count": len(rows), "format": req.format, "filename": f"defects.{suffix}", "content_type": content_type},
                    expires_at=datetime.now(timezone.utc) + timedelta(days=7), completed_at=datetime.now(timezone.utc))
    db.add(job)
    await db.commit()
    return success(_job_dict(job), "导出文件已生成")


def _job_dict(job):
    return {"id": str(job.id), "project_id": str(job.project_id), "kind": job.kind, "status": job.status,
            "summary": job.summary_json, "created_at": json_value(job.created_at), "completed_at": json_value(job.completed_at),
            "expires_at": json_value(job.expires_at), "download_url": f"/api/defect-management/jobs/{job.id}/download" if job.result_key else None}


@router.get("/jobs/{job_id}")
async def get_job(job_id: UUID, db=Depends(get_db_session), user=Depends(get_current_user)):
    job = await db.get(DefectJob, job_id)
    if not job:
        raise HTTPException(404, "作业不存在")
    context = await context_for(db, user, job.project_id)
    if job.creator_id != user.id and not context.privileged:
        raise HTTPException(404, "作业不存在")
    return success(_job_dict(job))


@router.get("/jobs/{job_id}/download")
async def download_job(job_id: UUID, db=Depends(get_db_session), user=Depends(get_current_user)):
    job = await db.get(DefectJob, job_id)
    if not job or not job.result_key or job.status != "completed":
        raise HTTPException(404, "下载文件不存在")
    context = await context_for(db, user, job.project_id)
    if job.creator_id != user.id and not context.privileged:
        raise HTTPException(404, "下载文件不存在")
    if job.expires_at and job.expires_at < datetime.now(timezone.utc):
        raise HTTPException(410, "下载文件已过期")
    response = await asyncio.to_thread(get_minio_client().get_object, settings.MINIO_BUCKET, job.result_key)
    def stream():
        try:
            for chunk in response.stream(64 * 1024):
                yield chunk
        finally:
            response.close(); response.release_conn()
    summary = job.summary_json or {}
    return StreamingResponse(stream(), media_type=summary.get("content_type", "application/octet-stream"),
        headers={"Content-Disposition": f"attachment; filename={summary.get('filename', 'defects.xlsx')}", "X-Content-Type-Options": "nosniff"})


@router.get("/metrics")
async def defect_metrics(project_id: UUID, days: int = Query(30, ge=7, le=365),
                         db=Depends(get_db_session), user=Depends(get_current_user)):
    await context_for(db, user, project_id)
    settings_row = await db.get(ProjectDefectSettings, project_id)
    tz_name = settings_row.timezone if settings_row else "Asia/Shanghai"
    try:
        tz = ZoneInfo(tz_name)
    except ZoneInfoNotFoundError:
        tz, tz_name = ZoneInfo("Asia/Shanghai"), "Asia/Shanghai"
    now = datetime.now(timezone.utc)
    start_local = (now.astimezone(tz).date() - timedelta(days=days - 1))
    start = datetime.combine(start_local, datetime.min.time(), tzinfo=tz).astimezone(timezone.utc)
    defects = (await db.execute(select(Defect).where(Defect.project_id == project_id))).scalars().all()
    events = (await db.execute(select(DefectEvent).where(DefectEvent.project_id == project_id,
        DefectEvent.occurred_at >= start).order_by(DefectEvent.occurred_at))).scalars().all()
    days_map = {start_local + timedelta(days=index): {"created": 0, "resolved": 0, "reopened": 0, "closed": 0}
                for index in range(days)}
    for defect in defects:
        created = defect.created_at.replace(tzinfo=timezone.utc) if defect.created_at and defect.created_at.tzinfo is None else defect.created_at
        if created and created >= start and created.astimezone(tz).date() in days_map:
            days_map[created.astimezone(tz).date()]["created"] += 1
    for event in events:
        key = event.occurred_at.astimezone(tz).date()
        if key not in days_map:
            continue
        if event.event_type == "verify_pass": days_map[key]["resolved"] += 1
        if event.event_type in {"reopen", "verify_fail"}: days_map[key]["reopened"] += 1
        if event.event_type == "close": days_map[key]["closed"] += 1
    active = [item for item in defects if not item.deleted_at]
    revisions = {item.workflow_revision_id for item in active if item.workflow_revision_id}
    definitions = {row.id: row.definition_json for row in (await db.execute(select(DefectWorkflowRevision).where(
        DefectWorkflowRevision.id.in_(revisions)))).scalars().all()} if revisions else {}
    backlog = []
    distributions = {"severity": {}, "priority": {}, "state": {}, "source": {}, "assignee": {}}
    for item in active:
        definition = definitions.get(item.workflow_revision_id, DEFAULT_WORKFLOW)
        state = next((value for value in definition.get("states", []) if value.get("key") == item.workflow_state), STATES.get(item.workflow_state, {}))
        if state.get("is_backlog"):
            backlog.append(item)
        for name, value in (("severity", item.severity.value), ("priority", item.priority), ("state", item.workflow_state),
                            ("source", item.source), ("assignee", str(item.assignee_id) if item.assignee_id else "unassigned")):
            distributions[name][value] = distributions[name].get(value, 0) + 1
    resolved_hours = [(item.first_resolved_at - (item.created_at.replace(tzinfo=timezone.utc) if item.created_at.tzinfo is None else item.created_at)).total_seconds() / 3600
                      for item in defects if item.first_resolved_at and item.created_at and item.first_resolved_at >= (item.created_at.replace(tzinfo=timezone.utc) if item.created_at.tzinfo is None else item.created_at)]
    first_events = {}
    for event in (await db.execute(select(DefectEvent).where(DefectEvent.project_id == project_id,
        DefectEvent.event_type != "created").order_by(DefectEvent.occurred_at))).scalars().all():
        first_events.setdefault(event.defect_id, event.occurred_at)
    response_hours = []
    for item in defects:
        first = first_events.get(item.id)
        created = item.created_at.replace(tzinfo=timezone.utc) if item.created_at and item.created_at.tzinfo is None else item.created_at
        if first and created and first >= created:
            response_hours.append((first - created).total_seconds() / 3600)
    legacy = any((item.source_metadata or {}).get("legacy_snapshot") for item in defects)
    return success({
        "period": {"days": days, "timezone": tz_name, "from": start.isoformat(), "to": now.isoformat(), "half_open": True},
        "summary": {"total": len(active), "backlog": len(backlog), "overdue": sum(bool(item.due_at and item.due_at < now) for item in backlog),
                    "resolved_in_period": sum(value["resolved"] for value in days_map.values()),
                    "reopened_in_period": sum(value["reopened"] for value in days_map.values()),
                    "mean_time_to_resolve_hours": round(mean(resolved_hours), 2) if resolved_hours else None,
                    "median_time_to_resolve_hours": round(median(resolved_hours), 2) if resolved_hours else None,
                    "mean_first_response_hours": round(mean(response_hours), 2) if response_hours else None,
                    "reopen_rate": round(sum(value["reopened"] for value in days_map.values()) / max(1, sum(value["resolved"] for value in days_map.values())), 4)},
        "trend": [{"date": key.isoformat(), **value} for key, value in sorted(days_map.items())],
        "distributions": distributions,
        "age_buckets": {"0_2d": sum((now - (item.created_at.replace(tzinfo=timezone.utc) if item.created_at.tzinfo is None else item.created_at)).days <= 2 for item in backlog),
                        "3_7d": sum(3 <= (now - (item.created_at.replace(tzinfo=timezone.utc) if item.created_at.tzinfo is None else item.created_at)).days <= 7 for item in backlog),
                        "8_30d": sum(8 <= (now - (item.created_at.replace(tzinfo=timezone.utc) if item.created_at.tzinfo is None else item.created_at)).days <= 30 for item in backlog),
                        "over_30d": sum((now - (item.created_at.replace(tzinfo=timezone.utc) if item.created_at.tzinfo is None else item.created_at)).days > 30 for item in backlog)},
        "data_quality": {"definition_version": 1, "history_complete": not legacy,
                         "message": "历史迁移缺陷不补造过程时效" if legacy else "指标基于缺陷事件与当前快照"},
    })


def _settings_dict(item):
    return {"project_id": str(item.project_id), "default_workflow_id": str(item.default_workflow_id) if item.default_workflow_id else None,
            "permission_config": item.permission_config or {}, "field_config": item.field_config or {},
            "notification_config": item.notification_config or {}, "timezone": item.timezone, "row_version": item.row_version}


@router.get("/settings")
async def get_settings(project_id: UUID, db=Depends(get_db_session), user=Depends(get_current_user)):
    await context_for(db, user, project_id)
    item = await db.get(ProjectDefectSettings, project_id)
    if item is None:
        item = ProjectDefectSettings(project_id=project_id)
        db.add(item); await db.flush()
    return success(_settings_dict(item))


@router.put("/settings")
async def update_settings(project_id: UUID, req: DefectSettingsInput, db=Depends(get_db_session), user=Depends(get_current_user)):
    context = await context_for(db, user, project_id, write=True)
    if not context.privileged:
        raise HTTPException(403, "缺陷配置仅限项目负责人维护")
    try: ZoneInfo(req.timezone)
    except ZoneInfoNotFoundError: raise HTTPException(422, "时区名称无效")
    allowed_fields = {"required_on_create", "required_on_triage", "attachment_limit_bytes"}
    allowed_permissions = {"allow_reporter_edit", "allow_developer_edit_assigned"}
    allowed_notifications = {"enabled_events", "notify_actor", "default_watchers"}
    if set(req.field_config) - allowed_fields or set(req.permission_config) - allowed_permissions or set(req.notification_config) - allowed_notifications:
        raise HTTPException(422, "配置包含不受支持的字段")
    configurable_fields = {"module_id", "iteration_id", "environment_profile_id", "reproduce_steps", "expected_result",
                           "actual_result", "assignee_id", "verifier_id", "affected_release_ids", "labels"}
    for key in ("required_on_create", "required_on_triage"):
        value = req.field_config.get(key, [])
        if not isinstance(value, list) or set(value) - configurable_fields:
            raise HTTPException(422, f"{key} 包含不可配置的字段")
    attachment_limit = req.field_config.get("attachment_limit_bytes", 200 * 1024 * 1024)
    if not isinstance(attachment_limit, int) or not 1024 <= attachment_limit <= 200 * 1024 * 1024:
        raise HTTPException(422, "附件总量上限必须在 1KB 到 200MB 之间")
    if any(not isinstance(value, bool) for value in req.permission_config.values()):
        raise HTTPException(422, "权限开关必须是布尔值")
    allowed_events = {"*", "created", "updated", "assigned", "comment_added", "comment_updated", "comment_deleted",
                      "attachment_added", "attachment_deleted", "triage", "start", "resume", "submit_fix",
                      "verify_pass", "verify_fail", "close", "reopen", "reject", "deleted", "restored", "workflow_migrated"}
    events = req.notification_config.get("enabled_events", ["*"])
    if not isinstance(events, list) or set(events) - allowed_events:
        raise HTTPException(422, "通知事件配置无效")
    if not isinstance(req.notification_config.get("notify_actor", False), bool):
        raise HTTPException(422, "通知操作者开关必须是布尔值")
    raw_watchers = req.notification_config.get("default_watchers", [])
    try:
        watcher_ids = {UUID(str(value)) for value in raw_watchers}
    except (ValueError, TypeError):
        raise HTTPException(422, "默认关注人编号无效")
    member_ids = set((await db.execute(select(ProjectMember.user_id).where(ProjectMember.project_id == project_id))).scalars().all())
    member_ids.add(context.project.owner_id)
    active_ids = set((await db.execute(select(User.id).where(User.id.in_(watcher_ids), User.is_active.is_(True)))).scalars().all()) if watcher_ids else set()
    if watcher_ids - (member_ids & active_ids):
        raise HTTPException(422, "默认关注人必须是当前项目有效成员")
    item = (await db.execute(select(ProjectDefectSettings).where(ProjectDefectSettings.project_id == project_id).with_for_update())).scalar_one_or_none()
    if item is None:
        item = ProjectDefectSettings(project_id=project_id); db.add(item); await db.flush()
    if item.row_version != req.expected_version:
        raise HTTPException(409, "项目缺陷配置已被修改，请刷新后重试")
    item.field_config, item.permission_config, item.notification_config, item.timezone = req.field_config, req.permission_config, req.notification_config, req.timezone
    item.row_version += 1
    await db.commit()
    return success(_settings_dict(item), "缺陷配置已保存")


def _validate_workflow(definition):
    if not isinstance(definition, dict):
        raise HTTPException(422, "工作流定义必须是对象")
    states = definition.get("states") or []
    transitions = definition.get("transitions") or []
    if not 2 <= len(states) <= 20:
        raise HTTPException(422, "工作流状态数必须在 2 到 20 之间")
    keys = [item.get("key") for item in states]
    if any(not isinstance(key, str) or not re.fullmatch(r"[a-z][a-z0-9_]{0,63}", key) for key in keys) or len(keys) != len(set(keys)):
        raise HTTPException(422, "状态标识必须唯一且使用小写字母、数字和下划线")
    if definition.get("initial_state") not in keys:
        raise HTTPException(422, "初始状态不存在")
    allowed_categories = {"todo", "in_progress", "verification", "done_verified", "done_rejected"}
    for state in states:
        if not state.get("label") or state.get("category") not in allowed_categories or state.get("legacy_status") not in {"open", "in_fix", "verified", "closed", "rejected"}:
            raise HTTPException(422, "状态标签、语义分类或兼容状态无效")
        if not isinstance(state.get("is_resolved"), bool) or not isinstance(state.get("is_backlog"), bool):
            raise HTTPException(422, "状态的解决/积压语义必须是布尔值")
    templates = {item["action_id"]: item for item in DEFAULT_WORKFLOW["transitions"]}
    actions = []
    for action in transitions:
        action_id = action.get("action_id")
        template = templates.get(action_id)
        if not template or action_id not in TRANSITION_FIELDS:
            raise HTTPException(422, "工作流只能使用平台支持的安全动作")
        if action_id in actions:
            raise HTTPException(422, "同一动作不能重复定义")
        actions.append(action_id)
        if not action.get("from") or any(value not in keys for value in action["from"]) or action.get("to") not in keys:
            raise HTTPException(422, f"动作 {action_id} 引用了不存在的状态")
        if not action.get("roles") or not action.get("actor_relations") or set(action.get("roles") or []) - ROLE_KEYS or set(action.get("actor_relations") or []) - set(template["actor_relations"]):
            raise HTTPException(422, f"动作 {action_id} 的权限范围无效")
        if action.get("required_fields", []) != template.get("required_fields", []) or action.get("required_any", []) != template.get("required_any", []):
            raise HTTPException(422, f"动作 {action_id} 不能改写平台必填规则")
        action["validators"], action["effects"], action["permission"] = template.get("validators", []), template.get("effects", []), template.get("permission")
    reachable = {definition["initial_state"]}
    changed = True
    while changed:
        changed = False
        for action in transitions:
            if reachable.intersection(action["from"]) and action["to"] not in reachable:
                reachable.add(action["to"]); changed = True
    if set(keys) - reachable:
        raise HTTPException(422, "存在从初始状态不可达的状态：" + "、".join(sorted(set(keys) - reachable)))
    definition["schema_version"] = 1
    definition["settings"] = {**DEFAULT_WORKFLOW["settings"], **{key: value for key, value in (definition.get("settings") or {}).items() if key == "allow_self_verification"}}
    definition["side_actions"] = deepcopy(DEFAULT_WORKFLOW.get("side_actions", []))
    definition["legacy_migration"] = {state["legacy_status"]: state["key"] for state in states}
    definition["rejection_resolutions"] = deepcopy(DEFAULT_WORKFLOW.get("rejection_resolutions", []))
    return definition


@router.get("/workflows")
async def list_workflows(project_id: UUID, db=Depends(get_db_session), user=Depends(get_current_user)):
    await context_for(db, user, project_id)
    workflows = (await db.execute(select(DefectWorkflow).where(DefectWorkflow.project_id == project_id).order_by(DefectWorkflow.name))).scalars().all()
    revisions = (await db.execute(select(DefectWorkflowRevision).where(DefectWorkflowRevision.project_id == project_id)
        .order_by(DefectWorkflowRevision.workflow_id, DefectWorkflowRevision.revision.desc()))).scalars().all()
    grouped = {}
    for revision in revisions:
        grouped.setdefault(revision.workflow_id, []).append({"id": str(revision.id), "revision": revision.revision,
            "status": revision.status, "published_at": json_value(revision.published_at)})
    return success({"items": [{"id": str(item.id), "name": item.name,
        "active_revision_id": str(item.active_revision_id) if item.active_revision_id else None,
        "revisions": grouped.get(item.id, [])} for item in workflows]})


@router.post("/workflows")
async def create_workflow(req: WorkflowCreate, db=Depends(get_db_session), user=Depends(get_current_user)):
    context = await context_for(db, user, req.project_id, write=True)
    if not context.privileged: raise HTTPException(403, "工作流仅限项目负责人维护")
    workflow = DefectWorkflow(id=uuid4(), project_id=req.project_id, name=req.name, created_by=user.id)
    definition = deepcopy(DEFAULT_WORKFLOW); definition["name"] = req.name; definition["revision"] = 1
    revision = DefectWorkflowRevision(id=uuid4(), workflow_id=workflow.id, project_id=req.project_id,
        revision=1, status="draft", definition_json=definition, published_by=None, published_at=None)
    db.add_all([workflow, revision])
    try: await db.commit()
    except IntegrityError: await db.rollback(); raise HTTPException(409, "同名工作流已存在")
    return success({"workflow_id": str(workflow.id), "revision_id": str(revision.id), "definition": definition})


@router.post("/workflows/{workflow_id}/draft")
async def create_draft(workflow_id: UUID, db=Depends(get_db_session), user=Depends(get_current_user)):
    workflow = await db.get(DefectWorkflow, workflow_id)
    if not workflow: raise HTTPException(404, "工作流不存在")
    context = await context_for(db, user, workflow.project_id, write=True)
    if not context.privileged: raise HTTPException(403, "工作流仅限项目负责人维护")
    existing = (await db.execute(select(DefectWorkflowRevision).where(DefectWorkflowRevision.workflow_id == workflow.id,
        DefectWorkflowRevision.status == "draft"))).scalar_one_or_none()
    if existing: return success({"id": str(existing.id), "revision": existing.revision, "definition": existing.definition_json})
    source = await db.get(DefectWorkflowRevision, workflow.active_revision_id) if workflow.active_revision_id else None
    if source is None:
        source = (await db.execute(select(DefectWorkflowRevision).where(DefectWorkflowRevision.workflow_id == workflow.id)
            .order_by(DefectWorkflowRevision.revision.desc()).limit(1))).scalar_one()
    number = (await db.execute(select(func.max(DefectWorkflowRevision.revision)).where(DefectWorkflowRevision.workflow_id == workflow.id))).scalar() or 0
    definition = deepcopy(source.definition_json); definition["revision"] = number + 1
    revision = DefectWorkflowRevision(id=uuid4(), workflow_id=workflow.id, project_id=workflow.project_id,
        revision=number + 1, status="draft", definition_json=definition, published_by=None, published_at=None)
    db.add(revision); await db.commit()
    return success({"id": str(revision.id), "revision": revision.revision, "definition": definition})


@router.put("/workflow-revisions/{revision_id}")
async def update_draft(revision_id: UUID, req: WorkflowDraftUpdate, db=Depends(get_db_session), user=Depends(get_current_user)):
    revision = (await db.execute(select(DefectWorkflowRevision).where(DefectWorkflowRevision.id == revision_id).with_for_update())).scalar_one_or_none()
    if not revision: raise HTTPException(404, "工作流版本不存在")
    context = await context_for(db, user, revision.project_id, write=True)
    if not context.privileged: raise HTTPException(403, "工作流仅限项目负责人维护")
    if revision.status != "draft": raise HTTPException(409, "已发布版本不可修改，请创建新草稿")
    definition = _validate_workflow(deepcopy(req.definition)); definition["revision"] = revision.revision
    revision.definition_json = definition
    await db.commit()
    return success({"id": str(revision.id), "definition": definition}, "草稿已保存")


@router.get("/workflow-revisions/{revision_id}")
async def get_workflow_revision(revision_id: UUID, db=Depends(get_db_session), user=Depends(get_current_user)):
    revision = await db.get(DefectWorkflowRevision, revision_id)
    if not revision:
        raise HTTPException(404, "工作流版本不存在")
    await context_for(db, user, revision.project_id)
    return success({"id": str(revision.id), "workflow_id": str(revision.workflow_id), "revision": revision.revision,
                    "status": revision.status, "definition": revision.definition_json,
                    "published_at": json_value(revision.published_at)})


@router.post("/workflow-revisions/{revision_id}/publish")
async def publish_revision(revision_id: UUID, db=Depends(get_db_session), user=Depends(get_current_user)):
    revision = (await db.execute(select(DefectWorkflowRevision).where(DefectWorkflowRevision.id == revision_id).with_for_update())).scalar_one_or_none()
    if not revision: raise HTTPException(404, "工作流版本不存在")
    context = await context_for(db, user, revision.project_id, write=True)
    if not context.privileged: raise HTTPException(403, "工作流仅限项目负责人维护")
    if revision.status != "draft": raise HTTPException(409, "该版本已经发布")
    revision.definition_json = _validate_workflow(deepcopy(revision.definition_json))
    revision.status, revision.published_by, revision.published_at = "published", user.id, datetime.now(timezone.utc)
    workflow = (await db.execute(select(DefectWorkflow).where(DefectWorkflow.id == revision.workflow_id).with_for_update())).scalar_one()
    workflow.active_revision_id, workflow.updated_at = revision.id, datetime.now(timezone.utc)
    settings_row = await db.get(ProjectDefectSettings, revision.project_id)
    if settings_row is None: settings_row = ProjectDefectSettings(project_id=revision.project_id); db.add(settings_row)
    settings_row.default_workflow_id = workflow.id; settings_row.row_version = int(settings_row.row_version or 0) + 1
    await db.commit()
    return success({"workflow_id": str(workflow.id), "active_revision_id": str(revision.id)}, "工作流已发布；仅新缺陷默认使用此版本")


@router.post("/workflow-migrations")
async def migrate_workflow(req: WorkflowMigrationRequest, project_id: UUID, db=Depends(get_db_session), user=Depends(get_current_user)):
    context = await context_for(db, user, project_id, write=True)
    if not context.privileged: raise HTTPException(403, "工作流迁移仅限项目负责人")
    digest = fingerprint(req.model_dump(mode="json", exclude={"preview"}))
    if not req.preview:
        prior = (await db.execute(select(DefectJob).where(DefectJob.project_id == project_id,
            DefectJob.creator_id == user.id, DefectJob.kind == "workflow_migration",
            DefectJob.request_key == req.request_key))).scalar_one_or_none()
        if prior:
            if prior.input_hash != digest:
                raise HTTPException(409, "同一请求编号不能用于不同迁移")
            return success({"job_id": str(prior.id), **(prior.summary_json or {})})
    target = await db.get(DefectWorkflowRevision, req.target_revision_id)
    if not target or target.project_id != project_id or target.status != "published": raise HTTPException(422, "目标工作流版本不可用")
    states = {item["key"] for item in target.definition_json["states"]}
    if set(req.state_mapping.values()) - states: raise HTTPException(422, "状态映射包含目标工作流不存在的状态")
    query = select(Defect).where(Defect.project_id == project_id, Defect.deleted_at.is_(None))
    if req.defect_ids: query = query.where(Defect.id.in_(req.defect_ids))
    defects = (await db.execute(query.order_by(Defect.id))).scalars().all()
    preview, errors = [], []
    for defect in defects:
        mapped = req.state_mapping.get(defect.workflow_state)
        if not mapped: errors.append({"id": str(defect.id), "state": defect.workflow_state, "message": "缺少状态映射"})
        else: preview.append({"id": str(defect.id), "defect_key": defect.defect_key, "from": defect.workflow_state, "to": mapped})
    if req.preview: return success({"items": preview, "errors": errors, "can_execute": not errors})
    if errors: raise HTTPException(422, {"message": "迁移预检失败", "errors": errors[:100]})
    for defect in defects:
        locked = (await db.execute(select(Defect).where(Defect.id == defect.id).with_for_update())).scalar_one()
        before = snapshot(locked); locked.workflow_state = req.state_mapping[locked.workflow_state]; locked.workflow_revision_id = target.id
        project_state(locked, target.definition_json); locked.row_version += 1; locked.updated_at = datetime.now(timezone.utc)
        append_event(db, locked, user, "workflow_migrated", before, extra={"target_revision_id": str(target.id)})
    job = DefectJob(id=uuid4(), project_id=project_id, creator_id=user.id, kind="workflow_migration", status="completed",
        request_key=req.request_key, input_hash=digest, summary_json={"migrated": len(defects)},
        completed_at=datetime.now(timezone.utc), expires_at=datetime.now(timezone.utc) + timedelta(days=30))
    db.add(job); await db.commit()
    return success({"job_id": str(job.id), "migrated": len(defects)}, "工作流迁移完成")
