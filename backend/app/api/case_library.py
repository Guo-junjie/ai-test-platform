"""
能力3（用例资产）API 路由 + 能力5/6/7 脚本绑定扩展

提供：
- POST /generate:       AI 生成并落库
- POST /:               人工录入手工用例
- GET /manual/options:  人工录入的目录、标签、需求选项
- GET  /:               列表
- GET  /{id}:           详情
- PUT  /{id}:           编辑
- DELETE /{id}:         删除
- POST /{id}/adopt:     单条接纳
- POST /{id}/deprecate: 单条废弃
- POST /adopt-batch:    批量接纳
- PUT  /{id}/scripts:   绑定脚本（能力5/6/7 扩展）
"""

import hashlib
import csv
import io
import json
import uuid
from datetime import datetime
from typing import Any

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field, ValidationError
from sqlalchemy import func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.schemas.case_library import (
    GenerateRequest,
    UpdateCaseRequest,
    AdoptBatchRequest,
    BatchUpdateCaseRequest,
    CaseAssetResponse,
    CaseImportCommitRequest,
    CaseImportRow,
    CreateManualCaseRequest,
)
from app.schemas.script import BindScriptRequest
from app.models.database import (
    ApiEndpoint,
    CaseAssetStatus,
    CaseReviewEvent,
    CaseReviewTask,
    Defect,
    DefectReference,
    InterfaceDoc,
    Project,
    RequirementDoc,
    TestCase,
    TestCaseAsset,
    TestCaseAssetVersion,
    TestPlan,
    TestPlanCase,
    TestPlanRevision,
    TestPlanRevisionCase,
    TestResult,
    TestRun,
    User,
    UserRole,
)
from app.modules.auth.dependencies import get_current_user, require_role
from app.modules.defects.workflow import STATE_LABELS
from app.modules.case_reviews import (
    cancel_pending_review_tasks,
    create_review_task,
    decide_review_task,
    reviewer_options,
    task_to_dict,
)
from app.modules.runs.case_readiness import api_case_errors
from app.utils.database import get_db_session

router = APIRouter()
case_editor = require_role(UserRole.SUPER_ADMIN, UserRole.ADMIN, UserRole.TEST_MANAGER, UserRole.TESTER)
case_reviewer = require_role(UserRole.SUPER_ADMIN, UserRole.ADMIN, UserRole.TEST_MANAGER)


class ReviewDecisionRequest(BaseModel):
    decision: str  # approve / changes_requested
    comment: str = Field(default="", max_length=2000)


def _enum_value(value: Any) -> Any:
    return getattr(value, "value", value)


def _normalize_tags(tags: list[str] | None) -> list[str]:
    result: list[str] = []
    for value in tags or []:
        tag = str(value).strip()[:40]
        if tag and tag not in result:
            result.append(tag)
    return result[:20]


def _case_snapshot(item: TestCaseAsset) -> dict[str, Any]:
    return {
        "title": item.title,
        "description": item.description or "",
        "execution_kind": item.execution_kind or "api",
        "case_type": item.case_type,
        "request_data": item.request_data or {},
        "expected_result": item.expected_result or {},
        "priority": item.priority or "P2",
        "folder": item.folder or "",
        "tags": list(item.tags or []),
        "pre_script": getattr(item, "pre_script", None) or "",
        "post_script": getattr(item, "post_script", None) or "",
        "sql_script": getattr(item, "sql_script", None) or "",
    }


def _add_case_version(
    db: AsyncSession,
    item: TestCaseAsset,
    actor_id: uuid.UUID | None,
    change_type: str,
    *,
    snapshot: dict[str, Any] | None = None,
    version: int | None = None,
) -> None:
    db.add(TestCaseAssetVersion(
        case_asset_id=item.id,
        version=version or int(item.version or 1),
        snapshot=snapshot or _case_snapshot(item),
        change_type=change_type,
        changed_by=actor_id,
    ))


def _case_to_dict(item: TestCaseAsset) -> dict[str, Any]:
    return {
        "id": str(item.id),
        "project_id": str(item.project_id),
        "endpoint_id": str(item.endpoint_id) if item.endpoint_id else None,
        "case_type": item.case_type,
        "title": item.title,
        "description": item.description or "",
        "request_data": item.request_data or {},
        "expected_result": item.expected_result,
        "priority": item.priority,
        "status": _enum_value(item.status),
        "execution_kind": item.execution_kind,
        "review_state": item.review_state,
        "source": _enum_value(item.source),
        "folder": item.folder or "",
        "tags": list(item.tags or []),
        "version": int(item.version or 1),
        "generation_context": item.generation_context or {},
        "created_by": str(item.created_by) if item.created_by else None,
        "created_at": item.created_at.isoformat() if item.created_at else None,
        "updated_at": item.updated_at.isoformat() if item.updated_at else None,
    }


def _record_review(db: AsyncSession, item: TestCaseAsset, actor: User,
                   action: str, comment: str = "") -> None:
    content = {"title": item.title, "description": item.description,
               "execution_kind": item.execution_kind, "case_type": item.case_type,
               "request_data": item.request_data, "expected_result": item.expected_result,
               "priority": item.priority, "pre_script": item.pre_script,
               "post_script": item.post_script, "sql_script": item.sql_script}
    fingerprint = hashlib.sha256(json.dumps(content, sort_keys=True, ensure_ascii=False,
                                             default=str).encode("utf-8")).hexdigest()
    db.add(CaseReviewEvent(case_asset_id=item.id, project_id=item.project_id,
                           actor_id=actor.id, action=action, comment=comment,
                           content_hash=fingerprint))


def _review_errors(item: TestCaseAsset) -> list[str]:
    errors = [] if (item.title or "").strip() else ["用例标题不能为空"]
    if item.execution_kind == "api":
        errors.extend(api_case_errors(item.request_data, item.expected_result))
    elif item.execution_kind == "manual":
        expected = item.expected_result if isinstance(item.expected_result, dict) else {}
        if not isinstance(expected.get("text"), str) or not expected["text"].strip():
            errors.append("手工用例需要明确预期结果")
    return errors


@router.post("/generate")
async def generate_cases(
    req: GenerateRequest,
    current_user: User = Depends(case_editor),
    db: AsyncSession = Depends(get_db_session),
) -> dict[str, Any]:
    """AI 生成测试用例并落库（DRAFT 状态）。

    支持三粒度：
      - 整项目（req.endpoint_ids / endpoint_id 均空）：自动取 project 下全部 active 接口
      - 多接口（req.endpoint_ids）：指定待生成的接口资产 id 列表
      - 单接口（req.endpoint_id）：指定单个接口资产 id
    """
    from app.models.database import TestCaseAsset, ApiEndpoint
    from app.modules.case_generator.case_generator import TestCaseGenerator

    # 校验 project_id 合法 UUID
    try:
        pid = uuid.UUID(req.project_id)
    except (ValueError, TypeError):
        raise HTTPException(status_code=400, detail="invalid project_id (must be UUID)")

    generator = TestCaseGenerator()

    # 构建 API 列表
    apis = []
    if req.endpoint_ids:
        # 多接口粒度
        for eid in req.endpoint_ids:
            try:
                ep_uuid = uuid.UUID(eid)
            except (ValueError, TypeError):
                continue
            result = await db.execute(
                select(ApiEndpoint).where(ApiEndpoint.id == ep_uuid)
            )
            ep = result.scalar_one_or_none()
            if ep and ep.project_id == pid:
                apis.append(_ep_to_dict(ep))
    elif req.endpoint_id:
        # 单接口粒度
        try:
            ep_uuid = uuid.UUID(req.endpoint_id)
        except (ValueError, TypeError):
            pass
        else:
            result = await db.execute(
                select(ApiEndpoint).where(ApiEndpoint.id == ep_uuid)
            )
            ep = result.scalar_one_or_none()
            if ep and ep.project_id == pid:
                apis.append(_ep_to_dict(ep))
    else:
        # 整项目粒度：自动取 project 下所有 active 接口（上限 30，避免一次生成过大）
        result = await db.execute(
            select(ApiEndpoint)
            .where(ApiEndpoint.project_id == pid, ApiEndpoint.is_active == True)
            .order_by(ApiEndpoint.method, ApiEndpoint.path)
            .limit(30)
        )
        for ep in result.scalars().all():
            apis.append(_ep_to_dict(ep))

    if not apis:
        raise HTTPException(
            status_code=400,
            detail="未找到可用接口：请确认项目下已有解析出的接口资产（数据源→接口文档），或在弹窗中勾选接口后再生成",
        )

    # 生成用例（能力12 P1：传项目 ID，知识库注入按项目过滤）
    cases = await generator.generate_all(apis, {}, project_id=str(pid))

    # 落库
    saved = []
    for case in cases.get("api", []):
        asset = TestCaseAsset(
            id=str(uuid.uuid4()),
            project_id=req.project_id,
            endpoint_id=(uuid.UUID(case["endpoint_id"]) if case.get("endpoint_id") else None),
            case_type=case.get("case_type", "positive"),
            # M2：语义拆分——AI 接口用例 execution_kind=api，设计类型沿用 case_type
            execution_kind="api",
            design_type=case.get("case_type", "positive"),
            title=case.get("case_name", ""),
            description=case.get("description", ""),
            request_data=case.get("request", {}),
            expected_result=case.get("expected"),
            priority=case.get("priority", "P2"),
            status="DRAFT",
            source="ai_generated",
            version=1,
            generation_context={
                "channel": "api_document",
                "endpoint_id": case.get("endpoint_id"),
                "generator": "ai_case_generator",
            },
            created_by=current_user.id,
            created_at=datetime.utcnow(),
            updated_at=datetime.utcnow(),
        )
        db.add(asset)
        saved.append(asset)

    await db.flush()
    for asset in saved:
        _add_case_version(db, asset, current_user.id, "generated")
    await db.flush()

    return {
        "code": 0,
        "data": {
            "cases": [
                {
                    "id": a.id,
                    "endpoint_id": a.endpoint_id,
                    "case_type": a.case_type,
                    "title": a.title,
                    "status": a.status,
                }
                for a in saved
            ],
            "total": len(saved),
            "inserted": len(saved),  # 兼容前端 caseApi.generate 返回字段（成功提示用）
        },
        "message": "ok",
    }


def _ep_to_dict(ep) -> dict[str, Any]:
    """ApiEndpoint ORM → 生成器所需 dict 的标准化映射（前后端字段名统一）。"""
    return {
        "endpoint_id": str(ep.id),
        "path": ep.path,
        "http_method": ep.method,  # ApiEndpoint 列名是 method，前端/生成器期望 http_method
        "params": ep.params or [],
        "auth_required": bool(ep.auth_required),
        "summary": ep.summary or "",
    }


@router.get("")
async def list_cases(
    project_id: str | None = None,
    endpoint_id: str | None = None,
    case_type: str | None = None,
    status: str | None = None,
    review_state: str | None = None,
    source: str | None = None,  # 新增：ai_generated / requirement / manual
    keyword: str | None = None,
    page: int = Query(1, ge=1),
    page_size: int = Query(20, ge=1, le=200),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db_session),
) -> dict[str, Any]:
    """获取用例列表。"""
    from app.models.database import TestCaseAsset

    query = select(TestCaseAsset).order_by(TestCaseAsset.created_at.desc())
    if project_id:
        query = query.where(TestCaseAsset.project_id == project_id)
    if endpoint_id:
        query = query.where(TestCaseAsset.endpoint_id == endpoint_id)
    if case_type:
        query = query.where(TestCaseAsset.case_type == case_type)
    if status:
        query = query.where(TestCaseAsset.status == status)
    if review_state:
        query = query.where(TestCaseAsset.review_state == review_state)
    if source:
        query = query.where(TestCaseAsset.source == source)
    if keyword:
        query = query.where(
            TestCaseAsset.title.ilike(f"%{keyword}%")
        )

    total = (await db.execute(select(func.count()).select_from(
        query.order_by(None).subquery()))).scalar_one()

    # Paginate
    offset = (page - 1) * page_size
    query = query.offset(offset).limit(page_size)
    result = await db.execute(query)
    items = result.scalars().all()

    return {
        "code": 0,
        "data": {
            "items": [_case_to_dict(item) for item in items],
            "total": total,
            "page": page,
            "page_size": page_size,
        },
        "message": "ok",
    }


@router.get("/manual/options")
async def manual_case_options(
    project_id: uuid.UUID,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db_session),
) -> dict[str, Any]:
    """人工录入所需的项目目录、标签和需求文档，不受列表分页限制。"""
    if not (await db.execute(select(Project.id).where(Project.id == project_id))).scalar_one_or_none():
        raise HTTPException(404, "项目不存在")
    rows = (await db.execute(select(TestCaseAsset.folder, TestCaseAsset.tags).where(
        TestCaseAsset.project_id == project_id,
    ))).all()
    documents = (await db.execute(select(RequirementDoc).where(
        RequirementDoc.project_id == project_id,
    ).order_by(RequirementDoc.created_at.desc()))).scalars().all()
    return {"code": 0, "data": {
        "folders": sorted({folder for folder, _ in rows if folder}),
        "tags": sorted({tag for _, tags in rows for tag in tags or []}),
        "requirements": [{"id": str(doc.id), "filename": doc.filename} for doc in documents],
    }, "message": "ok"}


@router.post("")
async def create_manual_case(
    req: CreateManualCaseRequest,
    current_user: User = Depends(case_editor),
    db: AsyncSession = Depends(get_db_session),
) -> dict[str, Any]:
    """创建人工用例及首版快照；同项目同标题需明确确认后才允许重复保存。"""
    project = (await db.execute(select(Project).where(
        Project.id == req.project_id,
    ).with_for_update())).scalar_one_or_none()
    if project is None:
        raise HTTPException(404, "项目不存在")
    if req.requirement_doc_id:
        document = (await db.execute(select(RequirementDoc).where(
            RequirementDoc.id == req.requirement_doc_id,
            RequirementDoc.project_id == req.project_id,
        ))).scalar_one_or_none()
        if document is None:
            raise HTTPException(422, "关联需求文档不存在或不属于当前项目")
    if not req.allow_duplicate:
        duplicates = (await db.execute(select(TestCaseAsset).where(
            TestCaseAsset.project_id == req.project_id,
            func.lower(func.trim(TestCaseAsset.title)) == req.title.lower(),
        ).order_by(TestCaseAsset.created_at.desc()).limit(10))).scalars().all()
        if duplicates:
            raise HTTPException(409, detail={
                "code": "CASE_DUPLICATE",
                "message": "当前项目已存在同名用例，请检查后确认是否继续添加",
                "duplicates": [{"id": str(item.id), "title": item.title,
                                "folder": item.folder or "未归类"} for item in duplicates],
            })
    request_data = {
        "preconditions": req.preconditions.strip(), "steps": req.steps,
        "requirement_ref": req.requirement_ref.strip(),
    }
    if req.requirement_doc_id:
        request_data["requirement_doc_id"] = str(req.requirement_doc_id)
    asset = TestCaseAsset(
        project_id=req.project_id, title=req.title, description=req.description.strip(),
        execution_kind="manual", case_type=req.case_type, design_type=req.case_type,
        priority=req.priority, folder=req.folder.strip() or None, tags=req.tags,
        request_data=request_data, expected_result={"text": req.expected_result},
        status=CaseAssetStatus.DRAFT, review_state="draft", source="manual", version=1,
        generation_context={"channel": "manual_entry"}, created_by=current_user.id,
    )
    db.add(asset)
    await db.flush()
    _add_case_version(db, asset, current_user.id, "created")
    await db.flush()
    return {"code": 0, "data": _case_to_dict(asset), "message": "用例已保存为草稿"}


@router.get("/{case_id}")
async def get_case(
    case_id: str,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db_session),
) -> dict[str, Any]:
    """获取用例详情。"""
    from app.models.database import TestCaseAsset

    result = await db.execute(
        select(TestCaseAsset).where(TestCaseAsset.id == case_id)
    )
    item = result.scalar_one_or_none()
    if item is None:
        raise HTTPException(status_code=404, detail="Case not found")

    return {
        "code": 0,
        "data": _case_to_dict(item),
        "message": "ok",
    }


@router.put("/{case_id}")
async def update_case(
    case_id: str,
    req: UpdateCaseRequest,
    current_user: User = Depends(case_editor),
    db: AsyncSession = Depends(get_db_session),
) -> dict[str, Any]:
    """编辑用例资产。"""
    from app.models.database import TestCaseAsset

    result = await db.execute(
        select(TestCaseAsset).where(TestCaseAsset.id == case_id).with_for_update()
    )
    item = result.scalar_one_or_none()
    if item is None:
        raise HTTPException(status_code=404, detail="Case not found")

    current_version = int(item.version or 1)
    prior_snapshot = _case_snapshot(item)
    has_current_version = (await db.execute(
        select(func.count()).select_from(TestCaseAssetVersion).where(
            TestCaseAssetVersion.case_asset_id == item.id,
            TestCaseAssetVersion.version == current_version,
        )
    )).scalar_one() > 0
    if not has_current_version:
        _add_case_version(
            db, item, item.created_by, "baseline",
            snapshot=prior_snapshot, version=current_version,
        )

    if item.review_state == "pending":
        await cancel_pending_review_tasks(
            db, item=item, actor_id=current_user.id, reason="用例内容被修改，原评审任务自动取消",
        )

    if req.title is not None:
        item.title = req.title
    if req.description is not None:
        item.description = req.description
    if req.request_data is not None:
        source_doc = (item.request_data or {}).get("requirement_doc_id")
        updated_request = dict(req.request_data)
        if source_doc:
            updated_request["requirement_doc_id"] = source_doc
        item.request_data = updated_request
    if req.expected_result is not None:
        item.expected_result = req.expected_result
    if req.priority is not None:
        item.priority = req.priority
    if req.case_type is not None:
        item.case_type = req.case_type
        item.design_type = req.case_type  # M2：设计类型同步
    if req.execution_kind is not None:
        if req.execution_kind not in {"manual", "api"}:
            raise HTTPException(422, "当前用例库只允许手工或 API 执行类型")
        item.execution_kind = req.execution_kind
    if req.folder is not None:
        item.folder = req.folder.strip() or None
    if req.tags is not None:
        item.tags = _normalize_tags(req.tags)

    if item.review_state in {"pending", "approved", "changes_requested"}:
        item.review_state = "draft"
        item.status = CaseAssetStatus.DRAFT
        _record_review(db, item, current_user, "edit", "内容变更，重新进入草稿")
    item.updated_at = datetime.utcnow()
    item.version = current_version + 1
    _add_case_version(db, item, current_user.id, "edit")
    await db.flush()

    return {"code": 0, "data": None, "message": "ok"}


@router.post("/batch-update")
async def batch_update_cases(
    req: BatchUpdateCaseRequest,
    current_user: User = Depends(case_editor),
    db: AsyncSession = Depends(get_db_session),
) -> dict[str, Any]:
    """批量归类和追加标签；已进入评审的用例修改后重新回到草稿。"""
    try:
        ids = list(dict.fromkeys(uuid.UUID(value) for value in req.ids))
    except (TypeError, ValueError) as exc:
        raise HTTPException(400, "用例 ID 格式错误") from exc
    items = (await db.execute(
        select(TestCaseAsset).where(TestCaseAsset.id.in_(ids)).with_for_update()
    )).scalars().all()
    if len(items) != len(ids):
        raise HTTPException(404, "部分用例不存在")

    existing_versions = {(row[0], row[1]) for row in (await db.execute(
        select(TestCaseAssetVersion.case_asset_id, TestCaseAssetVersion.version).where(
            TestCaseAssetVersion.case_asset_id.in_(ids)
        )
    )).all()}
    add_tags = _normalize_tags(req.add_tags)
    for item in items:
        current_version = int(item.version or 1)
        if (item.id, current_version) not in existing_versions:
            _add_case_version(db, item, item.created_by, "baseline", version=current_version)
        if item.review_state == "pending":
            await cancel_pending_review_tasks(
                db, item=item, actor_id=current_user.id, reason="用例归类或标签被修改，原评审任务自动取消",
            )
        if req.folder is not None:
            item.folder = req.folder.strip() or None
        item.tags = _normalize_tags([*(item.tags or []), *add_tags])
        item.version = current_version + 1
        item.updated_at = datetime.utcnow()
        if item.review_state in {"pending", "approved", "changes_requested"}:
            item.review_state = "draft"
            item.status = CaseAssetStatus.DRAFT
            _record_review(db, item, current_user, "edit", "归类或标签变更，重新进入草稿")
        _add_case_version(db, item, current_user.id, "classify")
    await db.flush()
    return {"code": 0, "data": {"count": len(items)}, "message": "批量更新完成"}


@router.post("/{case_id}/copy")
async def copy_case(
    case_id: str,
    current_user: User = Depends(case_editor),
    db: AsyncSession = Depends(get_db_session),
) -> dict[str, Any]:
    """复制为新的草稿用例，保留来源链路但重新进入评审。"""
    source = (await db.execute(
        select(TestCaseAsset).where(TestCaseAsset.id == case_id)
    )).scalar_one_or_none()
    if source is None:
        raise HTTPException(404, "用例不存在")
    generation_context = dict(source.generation_context or {})
    generation_context.update({"channel": "copy", "copied_from": str(source.id), "copied_from_version": int(source.version or 1)})
    copied = TestCaseAsset(
        project_id=source.project_id,
        endpoint_id=source.endpoint_id,
        case_type=source.case_type,
        execution_kind=source.execution_kind,
        design_type=source.design_type,
        title=f"{source.title}（副本）"[:500],
        description=source.description,
        request_data=dict(source.request_data or {}),
        expected_result=dict(source.expected_result or {}),
        priority=source.priority,
        folder=source.folder,
        tags=list(source.tags or []),
        version=1,
        generation_context=generation_context,
        status=CaseAssetStatus.DRAFT,
        review_state="draft",
        source="manual",
        created_by=current_user.id,
    )
    db.add(copied)
    await db.flush()
    _add_case_version(db, copied, current_user.id, "copied")
    await db.flush()
    return {"code": 0, "data": _case_to_dict(copied), "message": "用例已复制为草稿"}


@router.get("/{case_id}/versions")
async def list_case_versions(
    case_id: str,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db_session),
) -> dict[str, Any]:
    item = (await db.execute(select(TestCaseAsset).where(TestCaseAsset.id == case_id))).scalar_one_or_none()
    if item is None:
        raise HTTPException(404, "用例不存在")
    rows = (await db.execute(
        select(TestCaseAssetVersion, User.username)
        .outerjoin(User, User.id == TestCaseAssetVersion.changed_by)
        .where(TestCaseAssetVersion.case_asset_id == item.id)
        .order_by(TestCaseAssetVersion.version.desc())
    )).all()
    versions = [{
        "id": str(version.id),
        "version": version.version,
        "snapshot": version.snapshot or {},
        "change_type": version.change_type,
        "changed_by": str(version.changed_by) if version.changed_by else None,
        "changed_by_name": username,
        "created_at": version.created_at.isoformat() if version.created_at else None,
    } for version, username in rows]
    if not versions:
        versions.append({
            "id": None,
            "version": int(item.version or 1),
            "snapshot": _case_snapshot(item),
            "change_type": "baseline",
            "changed_by": str(item.created_by) if item.created_by else None,
            "changed_by_name": None,
            "created_at": item.created_at.isoformat() if item.created_at else None,
        })
    return {"code": 0, "data": {"items": versions}, "message": "ok"}


@router.get("/{case_id}/traceability")
async def get_case_traceability(
    case_id: str,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db_session),
) -> dict[str, Any]:
    item = (await db.execute(select(TestCaseAsset).where(TestCaseAsset.id == case_id))).scalar_one_or_none()
    if item is None:
        raise HTTPException(404, "用例不存在")

    plan_rows = (await db.execute(
        select(TestPlanCase, TestPlan)
        .join(TestPlan, TestPlan.id == TestPlanCase.plan_id)
        .where(TestPlanCase.case_asset_id == item.id)
        .order_by(TestPlan.updated_at.desc())
    )).all()
    revision_rows = (await db.execute(
        select(TestPlanRevisionCase, TestPlanRevision, TestPlan)
        .join(TestPlanRevision, TestPlanRevision.id == TestPlanRevisionCase.revision_id)
        .join(TestPlan, TestPlan.id == TestPlanRevision.plan_id)
        .where(TestPlanRevisionCase.case_asset_id == item.id)
        .order_by(TestPlanRevision.published_at.desc())
        .limit(50)
    )).all()
    run_rows = (await db.execute(
        select(TestCase, TestResult, TestRun)
        .join(TestRun, TestRun.id == TestCase.test_run_id)
        .outerjoin(TestResult, TestResult.test_case_id == TestCase.id)
        .where(TestCase.case_asset_id == item.id)
        .order_by(TestCase.created_at.desc())
        .limit(10)
    )).all()
    defect_rows = (await db.execute(
        select(Defect, TestCase)
        .outerjoin(TestCase, TestCase.id == Defect.test_case_id)
        .where(Defect.deleted_at.is_(None), or_(TestCase.case_asset_id == item.id,
            Defect.id.in_(select(DefectReference.defect_id).where(DefectReference.case_asset_id == item.id))))
        .order_by(Defect.created_at.desc())
        .limit(20)
    )).all()

    source_trace: dict[str, Any] = {
        "source": _enum_value(item.source),
        "generation_context": item.generation_context or {},
        "created_by": str(item.created_by) if item.created_by else None,
        "created_at": item.created_at.isoformat() if item.created_at else None,
    }
    if item.created_by:
        source_trace["created_by_name"] = (await db.execute(
            select(User.username).where(User.id == item.created_by)
        )).scalar_one_or_none()
    if item.endpoint_id:
        endpoint_row = (await db.execute(
            select(ApiEndpoint, InterfaceDoc)
            .outerjoin(InterfaceDoc, InterfaceDoc.id == ApiEndpoint.doc_id)
            .where(ApiEndpoint.id == item.endpoint_id)
        )).first()
        if endpoint_row:
            endpoint, document = endpoint_row
            source_trace["endpoint"] = {
                "id": str(endpoint.id), "method": endpoint.method, "path": endpoint.path,
                "version": endpoint.version, "document_id": str(document.id) if document else None,
                "document_name": document.filename if document else None,
            }
    requirement_doc_id = (item.request_data or {}).get("requirement_doc_id")
    if requirement_doc_id:
        requirement = (await db.execute(select(RequirementDoc).where(
            RequirementDoc.id == uuid.UUID(str(requirement_doc_id)), RequirementDoc.project_id == item.project_id,
        ))).scalar_one_or_none()
        if requirement:
            source_trace["requirement"] = {
                "id": str(requirement.id), "filename": requirement.filename,
                "parse_engine": requirement.parse_engine,
            }
    return {
        "code": 0,
        "data": {
            "source": source_trace,
            "plans": [{
                "id": str(plan.id), "name": plan.name, "status": plan.status,
                "enabled": bool(link.enabled), "added_at": link.added_at.isoformat() if link.added_at else None,
            } for link, plan in plan_rows],
            "plan_revisions": [{
                "plan_id": str(plan.id), "plan_name": plan.name,
                "revision": revision.revision, "revision_status": revision.status,
                "enabled": bool(link.enabled),
                "published_at": revision.published_at.isoformat() if revision.published_at else None,
            } for link, revision, plan in revision_rows],
            "recent_runs": [{
                "run_id": str(run.id), "run_status": _enum_value(run.status),
                "passed": result.is_passed if result else None,
                "status_code": result.status_code if result else None,
                "response_time_ms": result.response_time_ms if result else None,
                "executed_at": result.executed_at.isoformat() if result and result.executed_at else None,
            } for _case, result, run in run_rows],
            "defects": [{
                "id": str(defect.id), "title": defect.title,
                "severity": _enum_value(defect.severity),
                "status": STATE_LABELS.get(defect.workflow_state, defect.status),
                "workflow_state": defect.workflow_state,
                "created_at": defect.created_at.isoformat() if defect.created_at else None,
            } for defect, _case in defect_rows],
        },
        "message": "ok",
    }


def _first_import_value(raw: dict[str, Any], *keys: str) -> Any:
    for key in keys:
        value = raw.get(key)
        if value is not None and value != "":
            return value
    return None


def _json_import_cell(value: Any, default: Any, label: str, errors: list[str],
                      expected_type: type | tuple[type, ...] | None = None) -> Any:
    if value is None or value == "":
        return default
    parsed = value
    if isinstance(value, str):
        try:
            parsed = json.loads(value.strip())
        except (TypeError, ValueError, json.JSONDecodeError):
            errors.append(f"{label}必须是合法 JSON")
            return default
    if expected_type is not None and not isinstance(parsed, expected_type):
        errors.append(f"{label}格式不正确")
        return default
    return parsed


def _steps_import_cell(value: Any) -> list[str]:
    if isinstance(value, list):
        return [str(item).strip() for item in value]
    text = str(value or "").strip()
    if not text:
        return []
    try:
        parsed = json.loads(text)
        if isinstance(parsed, list):
            return [str(item).strip() for item in parsed]
    except (TypeError, ValueError, json.JSONDecodeError):
        pass
    separator = "||" if "||" in text else "\n"
    return [part.strip() for part in text.split(separator) if part.strip()]


def _body_import_cell(value: Any) -> Any:
    if value is None or value == "":
        return None
    if not isinstance(value, str):
        return value
    text = value.strip()
    try:
        return json.loads(text)
    except (TypeError, ValueError, json.JSONDecodeError):
        return text


def _validation_messages(exc: ValidationError) -> list[str]:
    messages: list[str] = []
    for error in exc.errors():
        message = str(error.get("msg") or "字段校验失败").removeprefix("Value error, ")
        if message not in messages:
            messages.append(message)
    return messages


def _normalize_import_row(raw: dict[str, Any], row_no: int) -> dict[str, Any]:
    errors: list[str] = []
    title = str(_first_import_value(raw, "title", "用例名称", "名称") or "").strip()
    kind_value = str(_first_import_value(raw, "execution_kind", "执行类型") or "api").strip().lower()
    kind = {"接口": "api", "接口用例": "api", "api用例": "api",
            "手工": "manual", "手工用例": "manual"}.get(kind_value, kind_value)
    case_type_value = str(_first_import_value(raw, "case_type", "设计类型", "用例类型") or
                          ("functional" if kind == "manual" else "positive")).strip().lower()
    case_type = {"功能": "functional", "正向": "positive", "反向": "negative",
                 "负向": "negative", "边界": "boundary", "异常": "exception"}.get(
                     case_type_value, case_type_value)
    priority = str(_first_import_value(raw, "priority", "优先级") or "P2").strip().upper()
    description = str(_first_import_value(raw, "description", "描述", "用例说明") or "")
    folder = str(_first_import_value(raw, "folder", "目录", "所属模块", "模块") or "").strip()
    tags_raw = _first_import_value(raw, "tags", "标签") or []
    tags = tags_raw if isinstance(tags_raw, list) else [value.strip() for value in
        str(tags_raw).replace("，", ",").split(",") if value.strip()]

    full_request = _first_import_value(raw, "request_data", "请求数据")
    request_data = _json_import_cell(full_request, {}, "请求数据", errors, dict) if full_request else {}
    raw_expected = _first_import_value(raw, "expected_result", "预期结果")

    if kind == "manual":
        steps_value = _first_import_value(raw, "steps", "操作步骤", "步骤")
        steps = _steps_import_cell(steps_value if steps_value is not None else request_data.get("steps"))
        request_data = {
            **request_data,
            "preconditions": str(_first_import_value(raw, "preconditions", "前置条件") or
                                 request_data.get("preconditions") or "").strip(),
            "steps": steps,
            "requirement_ref": str(_first_import_value(raw, "requirement_ref", "需求编号/链接",
                                                       "关联需求") or request_data.get("requirement_ref") or "").strip(),
        }
        requirement_doc_id = _first_import_value(raw, "requirement_doc_id", "关联需求文档ID") or request_data.get("requirement_doc_id")
        if requirement_doc_id:
            try:
                request_data["requirement_doc_id"] = str(uuid.UUID(str(requirement_doc_id)))
            except ValueError:
                errors.append("关联需求文档 ID 格式错误")
        if isinstance(raw_expected, dict):
            expected_result = raw_expected
        else:
            expected_text = str(raw_expected or "").strip()
            if expected_text.startswith("{"):
                parsed_expected = _json_import_cell(expected_text, {}, "预期结果", errors, dict)
                expected_result = parsed_expected
            else:
                expected_result = {"text": expected_text}
        if not steps:
            errors.append("手工用例至少需要一个操作步骤")
        if not isinstance(expected_result.get("text"), str) or not expected_result["text"].strip():
            errors.append("手工用例必须填写预期结果")
    else:
        method = str(_first_import_value(raw, "method", "请求方法") or request_data.get("method") or "GET").strip().upper()
        url = str(_first_import_value(raw, "url", "请求地址") or request_data.get("url") or "").strip()
        params_value = _first_import_value(raw, "params", "查询参数")
        headers_value = _first_import_value(raw, "headers", "请求头")
        request_data = {
            **request_data,
            "method": method,
            "url": url,
            "params": _json_import_cell(params_value, request_data.get("params", {}), "查询参数", errors, dict),
            "headers": _json_import_cell(headers_value, request_data.get("headers", {}), "请求头", errors, dict),
        }
        body_value = _first_import_value(raw, "body", "请求体")
        if body_value is not None:
            request_data["body"] = _body_import_cell(body_value)
        if isinstance(raw_expected, dict):
            expected_result = dict(raw_expected)
        elif raw_expected:
            expected_text = str(raw_expected).strip()
            if expected_text.startswith("{"):
                expected_result = _json_import_cell(expected_text, {}, "预期结果", errors, dict)
            else:
                expected_result = {"assertions": [{"type": "contains", "expected": expected_text}]}
        else:
            expected_result = {}
        status_value = _first_import_value(raw, "status_code", "期望状态码")
        if status_value is not None:
            try:
                expected_result["status_code"] = int(status_value)
            except (TypeError, ValueError):
                errors.append("期望状态码必须是整数")
        assertions_value = _first_import_value(raw, "assertions", "断言")
        if assertions_value is not None:
            expected_result["assertions"] = _json_import_cell(
                assertions_value, [], "断言", errors, list)

    data = {
        "title": title, "description": description, "execution_kind": kind,
        "case_type": case_type, "priority": priority, "folder": folder,
        "tags": tags, "request_data": request_data, "expected_result": expected_result,
    }
    try:
        validated = CaseImportRow.model_validate(data)
        data = validated.model_dump()
    except ValidationError as exc:
        errors.extend(_validation_messages(exc))
    if kind == "api":
        errors.extend(api_case_errors(request_data, expected_result))
    errors = list(dict.fromkeys(errors))
    return {"row_no": row_no, "valid": not errors, "errors": errors, "duplicate": False, "data": data}


@router.post("/import/preview")
async def preview_case_import(
    project_id: str = Form(...),
    file: UploadFile = File(...),
    current_user: User = Depends(case_editor),
    db: AsyncSession = Depends(get_db_session),
) -> dict[str, Any]:
    """解析 CSV/JSON/XLSX，并返回逐行错误；预检不写入数据库。"""
    try:
        project_uuid = uuid.UUID(project_id)
    except ValueError as exc:
        raise HTTPException(400, "项目 ID 格式错误") from exc
    if not (await db.execute(select(Project.id).where(Project.id == project_uuid))).scalar_one_or_none():
        raise HTTPException(404, "项目不存在")
    content = await file.read()
    suffix = (file.filename or "").lower().rsplit(".", 1)[-1]
    try:
        if suffix == "csv":
            text = content.decode("utf-8-sig")
            raw_rows = list(csv.DictReader(io.StringIO(text)))
        elif suffix == "json":
            parsed = json.loads(content.decode("utf-8-sig"))
            raw_rows = parsed if isinstance(parsed, list) else parsed.get("rows", []) if isinstance(parsed, dict) else []
        elif suffix in {"xlsx", "xlsm"}:
            from openpyxl import load_workbook
            sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).active
            values = list(sheet.iter_rows(values_only=True))
            headers = [str(v or "").strip() for v in (values[0] if values else [])]
            raw_rows = [dict(zip(headers, row)) for row in values[1:]]
        else:
            raise HTTPException(415, "仅支持 CSV、JSON、XLSX 文件")
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(422, f"文件解析失败：{exc}") from exc
    raw_rows = [row for row in raw_rows if isinstance(row, dict) and any(value not in (None, "") for value in row.values())]
    if not raw_rows:
        raise HTTPException(422, "文件中没有可识别的数据行")
    if len(raw_rows) > 2000:
        raise HTTPException(422, "单次最多导入 2000 条用例")
    rows = [_normalize_import_row(row, index + 2) for index, row in enumerate(raw_rows)]

    requirement_ids = {
        uuid.UUID(str(row["data"]["request_data"]["requirement_doc_id"]))
        for row in rows
        if row["data"].get("request_data", {}).get("requirement_doc_id")
        and "关联需求文档 ID 格式错误" not in row["errors"]
    }
    valid_requirement_ids = set((await db.execute(select(RequirementDoc.id).where(
        RequirementDoc.project_id == project_uuid,
        RequirementDoc.id.in_(requirement_ids),
    ))).scalars().all()) if requirement_ids else set()
    for row in rows:
        requirement_doc_id = row["data"].get("request_data", {}).get("requirement_doc_id")
        if requirement_doc_id and uuid.UUID(str(requirement_doc_id)) not in valid_requirement_ids:
            row["errors"].append("关联需求文档不存在或不属于当前项目")
            row["valid"] = False

    normalized_titles = {row["data"]["title"].strip().lower() for row in rows if row["data"].get("title")}
    existing_rows = (await db.execute(select(TestCaseAsset.id, TestCaseAsset.title, TestCaseAsset.folder).where(
        TestCaseAsset.project_id == project_uuid,
        func.lower(func.trim(TestCaseAsset.title)).in_(normalized_titles),
    ))).all() if normalized_titles else []
    existing_by_title: dict[str, list[dict[str, str]]] = {}
    for asset_id, title, folder in existing_rows:
        existing_by_title.setdefault(title.strip().lower(), []).append({
            "id": str(asset_id), "title": title, "folder": folder or "未归类",
        })
    seen_titles: set[str] = set()
    duplicate_count = 0
    for row in rows:
        key = row["data"].get("title", "").strip().lower()
        existing_duplicates = existing_by_title.get(key, [])
        row["duplicate"] = bool(key and (existing_duplicates or key in seen_titles))
        row["duplicates"] = existing_duplicates
        if row["duplicate"]:
            duplicate_count += 1
        if key:
            seen_titles.add(key)
    return {"code": 0, "data": {
        "project_id": project_id, "filename": file.filename, "rows": rows,
        "success_count": sum(1 for row in rows if row["valid"]),
        "failed_count": sum(1 for row in rows if not row["valid"]),
        "duplicate_count": duplicate_count,
    }, "message": "预检完成"}


@router.post("/import/commit")
async def commit_case_import(
    req: CaseImportCommitRequest,
    current_user: User = Depends(case_editor),
    db: AsyncSession = Depends(get_db_session),
) -> dict[str, Any]:
    """提交已通过预检的标准行，原子写入用例资产与首个版本。"""
    project_id = req.project_id
    if not (await db.execute(select(Project.id).where(Project.id == project_id))).scalar_one_or_none():
        raise HTTPException(404, "项目不存在")

    row_errors: list[str] = []
    for index, row in enumerate(req.rows, start=1):
        if row.execution_kind == "api":
            errors = api_case_errors(row.request_data, row.expected_result)
            if errors:
                row_errors.append(f"第 {index} 条：{'；'.join(errors)}")
    if row_errors:
        raise HTTPException(422, "导入数据未通过校验：" + "；".join(row_errors[:10]))

    try:
        requirement_ids = {
            uuid.UUID(str(row.request_data["requirement_doc_id"]))
            for row in req.rows if row.request_data.get("requirement_doc_id")
        }
    except ValueError as exc:
        raise HTTPException(422, "关联需求文档 ID 格式错误") from exc
    if requirement_ids:
        valid_requirement_ids = set((await db.execute(select(RequirementDoc.id).where(
            RequirementDoc.project_id == project_id,
            RequirementDoc.id.in_(requirement_ids),
        ))).scalars().all())
        if valid_requirement_ids != requirement_ids:
            raise HTTPException(422, "部分关联需求文档不存在或不属于当前项目")

    title_keys = {row.title.strip().lower() for row in req.rows}
    existing_titles = set((await db.execute(select(func.lower(func.trim(TestCaseAsset.title))).where(
        TestCaseAsset.project_id == project_id,
        func.lower(func.trim(TestCaseAsset.title)).in_(title_keys),
    ))).scalars().all()) if title_keys else set()
    assets: list[TestCaseAsset] = []
    skipped = 0
    accepted_titles: set[str] = set()
    for row in req.rows:
        data = row.model_dump()
        title_key = data["title"].strip().lower()
        if req.duplicate_strategy == "skip" and (title_key in existing_titles or title_key in accepted_titles):
            skipped += 1
            continue
        asset = TestCaseAsset(
            project_id=project_id, title=data["title"], description=data["description"],
            execution_kind=data["execution_kind"], case_type=data["case_type"], design_type=data["case_type"],
            priority=data["priority"], folder=data["folder"] or None,
            tags=_normalize_tags(data["tags"]), request_data=data["request_data"],
            expected_result=data["expected_result"], status=CaseAssetStatus.DRAFT,
            review_state="draft", source="manual", version=1,
            generation_context={"channel": "file_import"}, created_by=current_user.id,
        )
        db.add(asset)
        assets.append(asset)
        accepted_titles.add(title_key)
    if not assets:
        return {"code": 0, "data": {"imported": 0, "skipped": skipped, "asset_ids": []},
                "message": "没有需要导入的新用例"}
    await db.flush()
    for asset in assets:
        _add_case_version(db, asset, current_user.id, "imported")
    await db.flush()
    return {"code": 0, "data": {"imported": len(assets), "skipped": skipped,
                                   "asset_ids": [str(a.id) for a in assets]}, "message": "导入完成"}


@router.get("/import/template")
async def download_case_import_template(
    current_user: User = Depends(get_current_user),
):
    from openpyxl import Workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "用例导入"
    headers = [
        "用例名称", "描述", "执行类型", "设计类型", "优先级", "目录", "标签",
        "前置条件", "操作步骤", "关联需求文档ID", "需求编号/链接",
        "请求方法", "请求地址", "查询参数", "请求头", "请求体",
        "期望状态码", "断言", "预期结果",
    ]
    sheet.append(headers)
    sheet.append([
        "未登录用户无法提交订单", "验证订单权限控制", "manual", "negative", "P1", "订单/提交", "冒烟,权限",
        "购物车中已有商品", '["点击提交订单","观察页面提示"]', "", "REQ-102",
        "", "", "", "", "", "", "", "跳转登录页，且订单未创建",
    ])
    sheet.append([
        "查询商品列表", "正常查询商品", "api", "positive", "P1", "商品/查询", "冒烟,回归",
        "", "", "", "",
        "GET", "/api/products", "{}", "{}", "", "200", '[]', "",
    ])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:S{sheet.max_row}"
    for index, width in enumerate([28, 24, 13, 13, 10, 18, 18, 24, 36, 38, 24, 12, 28, 20, 20, 24, 13, 36, 32], start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = width
    guide = workbook.create_sheet("填写说明")
    guide.append(["字段", "说明"])
    guide_rows = [
        ("执行类型", "api 或 manual"),
        ("设计类型", "functional、positive、negative、boundary、exception"),
        ("操作步骤", "手工用例必填；推荐填写 JSON 数组，也支持用 || 或换行分隔"),
        ("预期结果", "手工用例填写可验证的文本；API 用例可填写响应包含文本"),
        ("期望状态码", "API 用例可填写 100–599；状态码或断言至少填写一项"),
        ("断言", "API 用例填写 JSON 数组，例如 [{\"type\":\"json_path\",\"path\":\"$.data.id\",\"operator\":\"not_null\"}]"),
        ("关联需求文档ID", "选填，必须是当前项目已有需求文档的 UUID"),
    ]
    for row in guide_rows:
        guide.append(row)
    guide.column_dimensions["A"].width = 22
    guide.column_dimensions["B"].width = 100
    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    headers = {"Content-Disposition": 'attachment; filename="case-import-template.xlsx"'}
    return StreamingResponse(output, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers=headers)


@router.delete("/{case_id}")
async def delete_case(
    case_id: str,
    current_user: User = Depends(require_role(UserRole.SUPER_ADMIN, UserRole.ADMIN, UserRole.TEST_MANAGER)),
    db: AsyncSession = Depends(get_db_session),
) -> dict[str, Any]:
    """删除用例资产。"""
    from app.models.database import TestCaseAsset

    result = await db.execute(
        select(TestCaseAsset).where(TestCaseAsset.id == case_id)
    )
    item = result.scalar_one_or_none()
    if item is None:
        raise HTTPException(status_code=404, detail="Case not found")

    if item.review_state != "draft" or item.status == CaseAssetStatus.ADOPTED:
        raise HTTPException(409, "已进入评审或计划流程的用例不能物理删除，请废弃")
    review_count = (await db.execute(select(func.count()).select_from(CaseReviewTask).where(
        CaseReviewTask.case_asset_id == item.id,
    ))).scalar_one()
    if review_count:
        raise HTTPException(409, "已有评审记录的用例不能物理删除，请废弃以保留审计链路")
    await db.delete(item)
    await db.flush()

    return {"code": 0, "data": None, "message": "ok"}


@router.post("/{case_id}/submit-review")
async def submit_case_review(
    case_id: str,
    current_user: User = Depends(case_editor),
    db: AsyncSession = Depends(get_db_session),
) -> dict[str, Any]:
    """兼容旧客户端：自动选择当前用户或项目首位可用评审人并创建正式任务。"""
    try:
        case_uuid = uuid.UUID(case_id)
    except ValueError as exc:
        raise HTTPException(400, "用例 ID 格式错误") from exc
    item = (await db.execute(select(TestCaseAsset).where(TestCaseAsset.id == case_uuid))).scalar_one_or_none()
    if item is None:
        raise HTTPException(404, "用例不存在")
    candidates = await reviewer_options(db, item.project_id)
    reviewer_id = next((uuid.UUID(row["id"]) for row in candidates if row["id"] == str(current_user.id)), None)
    reviewer_id = reviewer_id or (uuid.UUID(candidates[0]["id"]) if candidates else None)
    if reviewer_id is None:
        raise HTTPException(422, "当前项目没有可用评审人，请先配置项目成员")
    task = await create_review_task(
        db, case_id=case_uuid, reviewer_ids=[reviewer_id], comment="", due_at=None, actor=current_user,
    )
    return {"code": 0, "data": await task_to_dict(db, task, current_user.id), "message": "已创建评审任务"}


@router.post("/{case_id}/review")
async def review_case(
    case_id: str,
    req: ReviewDecisionRequest,
    current_user: User = Depends(case_reviewer),
    db: AsyncSession = Depends(get_db_session),
) -> dict[str, Any]:
    try:
        case_uuid = uuid.UUID(case_id)
    except ValueError as exc:
        raise HTTPException(400, "用例 ID 格式错误") from exc
    task = (await db.execute(select(CaseReviewTask).where(
        CaseReviewTask.case_asset_id == case_uuid,
        CaseReviewTask.status == "pending",
    ).order_by(CaseReviewTask.round.desc()).limit(1))).scalar_one_or_none()
    if task is None:
        raise HTTPException(409, "用例没有进行中的评审任务")
    task = await decide_review_task(
        db, task_id=task.id, decision=req.decision, comment=req.comment, actor=current_user,
    )
    return {"code": 0, "data": await task_to_dict(db, task, current_user.id), "message": "评审已记录"}


@router.get("/{case_id}/review-events")
async def get_case_review_events(
    case_id: str,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db_session),
) -> dict[str, Any]:
    item = (await db.execute(select(TestCaseAsset.id).where(TestCaseAsset.id == case_id))).scalar_one_or_none()
    if item is None:
        raise HTTPException(404, "用例不存在")
    rows = (await db.execute(select(CaseReviewEvent, User.username)
                             .join(User, User.id == CaseReviewEvent.actor_id)
                             .where(CaseReviewEvent.case_asset_id == item)
                             .order_by(CaseReviewEvent.created_at.asc()))).all()
    return {"code": 0, "data": [{"id": str(row.id), "actor_id": str(row.actor_id),
                                   "actor_name": username, "action": row.action,
                                   "comment": row.comment, "content_hash": row.content_hash,
                                   "task_id": str(row.task_id) if row.task_id else None,
                                   "case_version_id": str(row.case_version_id) if row.case_version_id else None,
                                   "case_version": row.case_version,
                                   "created_at": row.created_at.isoformat() if row.created_at else None}
                                  for row, username in rows], "message": "ok"}


@router.post("/{case_id}/adopt")
async def adopt_case(
    case_id: str,
    current_user: User = Depends(require_role(UserRole.SUPER_ADMIN, UserRole.ADMIN, UserRole.TEST_MANAGER)),
    db: AsyncSession = Depends(get_db_session),
) -> dict[str, Any]:
    """兼容旧客户端：把“接纳”映射为当前评审任务的通过结论。"""
    try:
        case_uuid = uuid.UUID(case_id)
    except ValueError as exc:
        raise HTTPException(400, "用例 ID 格式错误") from exc
    task = (await db.execute(select(CaseReviewTask).where(
        CaseReviewTask.case_asset_id == case_uuid,
        CaseReviewTask.status == "pending",
    ).order_by(CaseReviewTask.round.desc()).limit(1))).scalar_one_or_none()
    if task is None:
        raise HTTPException(409, "请先创建评审任务")
    task = await decide_review_task(
        db, task_id=task.id, decision="approve", comment="", actor=current_user,
    )
    return {"code": 0, "data": await task_to_dict(db, task, current_user.id), "message": "评审已通过"}


@router.post("/{case_id}/deprecate")
async def deprecate_case(
    case_id: str,
    current_user: User = Depends(require_role(UserRole.SUPER_ADMIN, UserRole.ADMIN, UserRole.TEST_MANAGER)),
    db: AsyncSession = Depends(get_db_session),
) -> dict[str, Any]:
    """单条废弃用例。"""
    from app.models.database import TestCaseAsset

    result = await db.execute(
        select(TestCaseAsset).where(TestCaseAsset.id == case_id)
    )
    item = result.scalar_one_or_none()
    if item is None:
        raise HTTPException(status_code=404, detail="Case not found")

    if item.review_state == "pending":
        await cancel_pending_review_tasks(
            db, item=item, actor_id=current_user.id, reason="用例被废弃，评审任务自动取消",
        )
    item.status = CaseAssetStatus.DEPRECATED
    item.review_state = "changes_requested"
    item.updated_at = datetime.utcnow()
    _record_review(db, item, current_user, "deprecate")
    await db.flush()

    return {"code": 0, "data": None, "message": "ok"}


@router.post("/{case_id}/restore")
async def restore_case(
    case_id: str,
    current_user: User = Depends(require_role(UserRole.SUPER_ADMIN, UserRole.ADMIN, UserRole.TEST_MANAGER)),
    db: AsyncSession = Depends(get_db_session),
) -> dict[str, Any]:
    """将归档用例恢复为草稿，恢复后必须重新提交评审。"""
    item = (await db.execute(
        select(TestCaseAsset).where(TestCaseAsset.id == case_id).with_for_update()
    )).scalar_one_or_none()
    if item is None:
        raise HTTPException(404, "用例不存在")
    if item.status != CaseAssetStatus.DEPRECATED:
        raise HTTPException(409, "只有已归档用例可以恢复")
    item.status = CaseAssetStatus.DRAFT
    item.review_state = "draft"
    item.updated_at = datetime.utcnow()
    _record_review(db, item, current_user, "restore", "从归档恢复为草稿")
    await db.flush()
    return {"code": 0, "data": {"review_state": "draft"}, "message": "已恢复为草稿"}


@router.post("/adopt-batch")
async def adopt_batch(
    req: AdoptBatchRequest,
    current_user: User = Depends(require_role(UserRole.SUPER_ADMIN, UserRole.ADMIN, UserRole.TEST_MANAGER)),
    db: AsyncSession = Depends(get_db_session),
) -> dict[str, Any]:
    """兼容旧客户端：批量提交当前用户被指派任务的通过结论。"""
    try:
        ids = list(dict.fromkeys(uuid.UUID(value) for value in req.ids))
    except (TypeError, ValueError) as exc:
        raise HTTPException(400, "用例 ID 格式错误") from exc
    if not ids:
        raise HTTPException(422, "至少选择一条用例")
    tasks = (await db.execute(select(CaseReviewTask).where(
        CaseReviewTask.case_asset_id.in_(ids),
        CaseReviewTask.status == "pending",
    ))).scalars().all()
    task_by_case = {task.case_asset_id: task for task in tasks}
    if len(task_by_case) != len(ids):
        raise HTTPException(409, "部分用例没有进行中的评审任务")
    for case_id in ids:
        await decide_review_task(
            db, task_id=task_by_case[case_id].id, decision="approve", comment="", actor=current_user,
        )
    return {"code": 0, "data": {"count": len(ids)}, "message": "批量评审结论已记录"}


# ============ 能力5/6/7 扩展：脚本绑定 ============

@router.put("/{case_id}/scripts")
async def bind_scripts(
    case_id: str,
    req: BindScriptRequest,
    current_user: User = Depends(case_editor),
    db: AsyncSession = Depends(get_db_session),
) -> dict[str, Any]:
    """
    绑定脚本到用例资产（能力5/6/7 扩展）。

    将生成的前置/后置/SQL 脚本绑定到 TestCaseAsset 的 pre_script/post_script/sql_script 字段。
    """
    from app.models.database import TestCaseAsset

    result = await db.execute(
        select(TestCaseAsset).where(TestCaseAsset.id == case_id).with_for_update()
    )
    item = result.scalar_one_or_none()
    if item is None:
        raise HTTPException(status_code=404, detail="Case not found")

    current_version = int(item.version or 1)
    prior_snapshot = _case_snapshot(item)
    has_current_version = (await db.execute(
        select(func.count()).select_from(TestCaseAssetVersion).where(
            TestCaseAssetVersion.case_asset_id == item.id,
            TestCaseAssetVersion.version == current_version,
        )
    )).scalar_one() > 0
    if not has_current_version:
        _add_case_version(db, item, item.created_by, "baseline", snapshot=prior_snapshot, version=current_version)
    if item.review_state == "pending":
        await cancel_pending_review_tasks(
            db, item=item, actor_id=current_user.id, reason="用例脚本被修改，原评审任务自动取消",
        )

    if req.pre_script is not None:
        item.pre_script = req.pre_script
    if req.post_script is not None:
        item.post_script = req.post_script
    if req.sql_script is not None:
        item.sql_script = req.sql_script

    if item.review_state in {"pending", "approved", "changes_requested"}:
        item.review_state = "draft"
        item.status = CaseAssetStatus.DRAFT
        _record_review(db, item, current_user, "edit", "脚本变更，重新进入草稿")
    item.updated_at = datetime.utcnow()
    item.version = current_version + 1
    _add_case_version(db, item, current_user.id, "script_edit")
    await db.flush()

    return {
        "code": 0,
        "data": {
            "id": item.id,
            "has_pre_script": bool(item.pre_script),
            "has_post_script": bool(item.post_script),
            "has_sql_script": bool(item.sql_script),
        },
        "message": "ok",
    }
